from __future__ import annotations

"""Least-squares refinement for the EDM/TDM chromatography models.

Two objective modes are intentionally supported:
1. CHROMATOGRAM: ordinary point-by-point least-squares against an attached raw
   chromatogram CSV.
2. CHROMATOGRAM_AND_COMPOSITION: the same chromatogram objective plus
   species mass-percent observations integrated over fraction intervals (CV or mL).

UV predictions use the target and species response factors; only eligible
unlocked factors are refined. Fit entry points require a UV reference in mAU
(AU-labelled columns are converted). RAW_SSE defaults to literal
pointwise equation-9 residuals without inferred baseline or normalization.
NORMALIZED_MSE and initial-median baseline correction are explicit options.
Optional composition observations add normalized composition residuals.
"""

from dataclasses import dataclass
import codecs
import copy
import csv
import html
import io
import json
import math
import os
from pathlib import Path
import re
import shutil
import tempfile
import time
from types import SimpleNamespace
from typing import Any, Iterable

import numpy as np
from scipy.interpolate import PchipInterpolator
from scipy.optimize import least_squares

from tdm_minimal_io import (
    CONFIG_PATH,
    MODEL_EDM_CPA,
    MODEL_EDM_LANGMUIR,
    MODEL_TDM_CPA,
    MODEL_TDM_LANGMUIR,
    MODEL_TDM_WANG,
    MAX_LSQ_PROFILES,
    active_components,
    canonical_fit_parameter_path,
    normalize_config,
    required_parameter_paths,
    save_config,
    selected_components,
    validate_config,
    is_hic,
)
from tdm_minimal_model import simulate, _build_program, _svg_plot, _stage_receipt

PROJECT_DIR = Path(__file__).resolve().parent
ATTACHED_CHROMATOGRAM_CSV = PROJECT_DIR / "least_squares_reference_chromatogram.csv"
FIT_RESULTS_CSV = PROJECT_DIR / "least_squares_fit_results.csv"
FIT_TRACE_CSV = PROJECT_DIR / "least_squares_fit_trace.csv"
FIT_CONDUCTIVITY_CSV = PROJECT_DIR / "least_squares_fit_conductivity.csv"
FIT_COMPOSITION_CSV = PROJECT_DIR / "least_squares_fit_composition.csv"
FIT_REPORT_HTML = PROJECT_DIR / "least_squares_fit_report.html"
LSQ_ALIGNMENT_HTML = PROJECT_DIR / "least_squares_alignment_preview.html"
FIT_CONFIG_JSON = PROJECT_DIR / "least_squares_fitted_inputs.json"
FIT_APPLY_JSL = PROJECT_DIR / "least_squares_apply_to_ui.jsl"
FIT_SETTINGS_JSON = PROJECT_DIR / "least_squares_settings.json"
FIT_HISTORY_CSV = PROJECT_DIR / "least_squares_objective_history.csv"
FIT_IMPROVED_CSV = PROJECT_DIR / "least_squares_improved_parameters.csv"
FIT_IMPROVED_TXT = PROJECT_DIR / "least_squares_improved_parameters.txt"

MODE_CHROMATOGRAM = "CHROMATOGRAM"
MODE_CHROMATOGRAM_AND_COMPOSITION = "CHROMATOGRAM_AND_COMPOSITION"
SUPPORTED_MODES = {MODE_CHROMATOGRAM, MODE_CHROMATOGRAM_AND_COMPOSITION}
MAX_COMPOSITION_POINTS = 30
SAMPLE_ORDER_FIELD = "__sample_order__"


@dataclass(frozen=True)
class ChromatogramReference:
    cv: np.ndarray
    signal: np.ndarray
    x_column: str
    signal_column: str
    x_unit: str
    source_path: str
    x_axis_assumption: str | None = None
    point_weights: np.ndarray | None = None
    weight_column: str | None = None
    conductivity_cv: np.ndarray | None = None
    conductivity_signal: np.ndarray | None = None
    conductivity_column: str | None = None
    sheet_name: str | None = None


@dataclass(frozen=True)
class FitParameter:
    path: str
    label: str
    transform: str  # LOG10 or LINEAR
    lower: float
    upper: float
    initial: float


def _norm_header(text: str) -> str:
    return "".join(ch for ch in str(text).strip().lower() if ch.isalnum())


def _parse_numeric_cell(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip().strip('"').strip("'")
    if not text or text.lower() in {"na", "n/a", "nan", "null", "none", "--", "-"}:
        return None
    text = text.replace("\u00a0", "").replace("\u202f", "").replace(" ", "").replace("−", "-")
    # Instrument exports often store elapsed time as mm:ss or hh:mm:ss.
    duration = re.fullmatch(r"([+-]?\d+):([0-5]?\d)(?::([0-5]?\d(?:[.,]\d+)?))?", text)
    if duration:
        first = float(duration.group(1))
        second = float(duration.group(2))
        third = duration.group(3)
        if third is None:
            return first * 60.0 + second
        return first * 3600.0 + second * 60.0 + float(third.replace(",", "."))
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        left, right = text.rsplit(",", 1)
        if right.isdigit() and 1 <= len(right) <= 6:
            text = left.replace(",", "") + "." + right
        else:
            text = text.replace(",", "")
    try:
        out = float(text)
    except Exception:
        return None
    return out if math.isfinite(out) else None


def _is_number(value: Any) -> bool:
    return _parse_numeric_cell(value) is not None


def _parse_path(path: str) -> list[str | int]:
    import re
    out: list[str | int] = []
    for name, index in re.findall(r"([^\.\[\]]+)|\[(\d+)\]", path):
        out.append(name if name else int(index))
    return out


def _get_path(data: Any, path: str) -> Any:
    cur = data
    for token in _parse_path(path):
        cur = cur[token]
    return cur


def _set_path(data: Any, path: str, value: Any) -> None:
    tokens = _parse_path(path)
    cur = data
    for token in tokens[:-1]:
        cur = cur[token]
    cur[tokens[-1]] = value


def _normalize_chromatogram_path_text(path: Path | str) -> str:
    """Clean file-picker text, including JMP's leading slash before Windows drives."""
    source_text = str(path).strip().strip('"').strip("'")
    if not source_text:
        raise ValueError("Chromatogram file path cannot be blank.")
    # JMP may return its POSIX-style Windows path (for example, /C:/data/run.csv)
    # or a slash-converted variant (\\C:\\data\\run.csv). Strip only the
    # separator(s) immediately before a drive letter; preserve UNC paths.
    return re.sub(r"^[\\/]+(?=[A-Za-z]:[\\/])", "", source_text)


def _resolve_chromatogram_path(path: Path | str) -> Path:
    """Resolve a chromatogram CSV/Excel file, including JMP path variants."""
    received_text = str(path).strip().strip('"').strip("'")
    source_text = _normalize_chromatogram_path_text(path)
    candidate = Path(source_text).expanduser()
    if candidate.is_file():
        return candidate.resolve()
    if not candidate.suffix:
        for extension in (".csv", ".CSV", ".xlsx", ".xlsm"):
            with_extension = Path(str(candidate) + extension)
            if with_extension.is_file():
                return with_extension.resolve()
        raise FileNotFoundError(
            "Chromatogram CSV/Excel file not found. Checked the entered path and "
            f"common extensions: {candidate}. In JMP, click Attach / validate chromatogram "
            "to choose the file."
        )
    raise FileNotFoundError(
        "Chromatogram CSV not found after JMP path normalization. "
        f"Path received: {received_text!r}. Path checked: {candidate}. "
        "Check the full filename and extension, or choose the file with the JMP picker."
    )


def _read_numeric_rows(path: Path, *, sheet_name: str | None = None) -> tuple[list[str], list[dict[str, str]], str | None]:
    if not path.is_file():
        raise FileNotFoundError(f"Chromatogram file not found: {path}")
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        # Keep the Excel dependency optional for users fitting CSV exports only.
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Excel chromatograms require openpyxl (pip install -r requirements.txt).") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            best = None
            sheets = workbook.worksheets
            if sheet_name:
                sheets = [sheet for sheet in sheets if sheet.title.strip().casefold() == sheet_name.strip().casefold()]
                if not sheets:
                    raise ValueError(f"Excel worksheet {sheet_name!r} not found. Available sheets: {', '.join(workbook.sheetnames)}")
            for sheet in sheets:
                records = []
                for row in sheet.iter_rows(values_only=True):
                    if any(cell is not None for cell in row):
                        records.append([str(c) if c is not None else "" for c in row])
                if not records:
                    continue
                buffer = io.StringIO()
                csv.writer(buffer).writerows(records)
                try:
                    fields, rows = _read_numeric_rows_text(buffer.getvalue())
                    x, y, _unit = _detect_columns(fields, rows)
                    count = sum(_is_number(r.get(y)) and (x == SAMPLE_ORDER_FIELD or _is_number(r.get(x))) for r in rows)
                except ValueError:
                    continue
                if best is None or count > best[0]:
                    best = (count, fields, rows, sheet.title)
            if best is None:
                raise ValueError("No worksheet has at least three numeric chromatogram samples; use headered CV/time and signal columns.")
            return best[1], best[2], best[3]
        finally:
            workbook.close()
    if path.suffix.lower() == ".xls":
        raise ValueError("Legacy .xls is not supported. Save the Excel workbook as .xlsx or CSV.")
    if sheet_name:
        raise ValueError("Worksheet selection can only be used with Excel .xlsx/.xlsm files, not CSV.")
    with path.open("rb") as fh:
        prefix = fh.read(4)
    if prefix.startswith(codecs.BOM_UTF32_LE) or prefix.startswith(codecs.BOM_UTF32_BE):
        encodings = ("utf-32",)
    elif prefix.startswith(codecs.BOM_UTF16_LE) or prefix.startswith(codecs.BOM_UTF16_BE):
        encodings = ("utf-16",)
    elif prefix.startswith(codecs.BOM_UTF8):
        encodings = ("utf-8-sig",)
    else:
        # UTF-8 is preferred. CSVs exported by Windows applications may instead
        # use the active Windows code page, so retry those common encodings.
        encodings = ("utf-8-sig", "cp1252", "latin-1")

    last_decode_error: UnicodeDecodeError | None = None
    for encoding in encodings:
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError as exc:
            last_decode_error = exc
    else:
        raise ValueError(
            "Chromatogram CSV text encoding could not be decoded. Supported formats "
            "include UTF-8, UTF-16/UTF-32 with a byte-order mark, and Windows-1252."
        ) from last_decode_error

    fields, rows = _read_numeric_rows_text(text)
    return fields, rows, None


def _read_numeric_rows_text(text: str) -> tuple[list[str], list[dict[str, str]]]:
    delimiters: list[str] = []
    try:
        guessed = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
        delimiters.append(guessed)
    except csv.Error:
        pass
    delimiters.extend(delim for delim in (",", ";", "\t", "|") if delim not in delimiters)

    best: tuple[tuple[int, int, int, int], list[str], list[dict[str, str]]] | None = None
    for delimiter in delimiters:
        try:
            records = list(csv.reader(io.StringIO(text), delimiter=delimiter, skipinitialspace=True))
        except csv.Error:
            continue
        records = [[str(cell).strip().lstrip("\ufeff") for cell in record] for record in records]
        records = [record for record in records if any(cell for cell in record)]
        for header_index, raw_header in enumerate(records[:40]):
            if not raw_header or not any(raw_header):
                continue
            fields = [cell or f"Column_{i + 1}" for i, cell in enumerate(raw_header)]
            # Make repeated/blank instrument column names safe for DictReader-like use.
            seen: dict[str, int] = {}
            unique_fields: list[str] = []
            for i, field in enumerate(fields):
                count = seen.get(field, 0) + 1
                seen[field] = count
                unique_fields.append(field if count == 1 else f"{field}_{count}")
            data_records = records[header_index + 1:]
            pair_rows = 0
            numeric_cells = 0
            single_numeric_rows = 0
            for record in data_records:
                values = [_parse_numeric_cell(record[i] if i < len(record) else "") for i in range(len(fields))]
                count = sum(value is not None for value in values)
                numeric_cells += count
                if count >= 2:
                    pair_rows += 1
                elif count == 1:
                    single_numeric_rows += 1
            if len(fields) > 1:
                usable_rows = pair_rows
            else:
                usable_rows = single_numeric_rows
            if usable_rows < 3:
                continue
            normalized_header = " ".join(_norm_header(field) for field in unique_fields)
            header_quality = sum(
                token in normalized_header
                for token in ("time", "volume", "cv", "retention", "sample", "index", "uv", "signal", "absorbance", "detector", "response", "concentration")
            )
            score = (usable_rows, numeric_cells, header_quality, header_index)
            data: list[dict[str, str]] = []
            for record in data_records:
                data.append({field: record[i].strip() if i < len(record) else "" for i, field in enumerate(unique_fields)})
            if best is None or score > best[0]:
                best = (score, unique_fields, data)

        # Headerless numeric exports are common in detector-only files. Preserve
        # the samples and synthesize column names if the first record is numeric.
        if records and all(_parse_numeric_cell(cell) is not None for cell in records[0]):
            width = len(records[0])
            fields = [f"Column_{i + 1}" for i in range(width)]
            data = [{field: record[i].strip() if i < len(record) else "" for i, field in enumerate(fields)} for record in records]
            usable = sum(sum(_parse_numeric_cell(row[field]) is not None for field in fields) >= min(2, width) for row in data)
            if usable >= 3:
                score = (usable, usable * width, 0, -1)
                if best is None or score > best[0]:
                    best = (score, fields, data)

    if best is None:
        raise ValueError(
            "Could not read at least three chromatogram samples from this CSV. Tried comma, "
            "semicolon, tab, and pipe separators, including instrument metadata rows. "
            "Export a table with a header and numeric detector values; a numeric time/CV/volume "
            "column is preferred, but a single detector column can use sample order."
        )
    _score, fields, rows = best
    if len(rows) < 3:
        raise ValueError("Chromatogram CSV must contain at least three data rows.")
    return fields, rows


def _column_numeric_count(rows: list[dict[str, str]], field: str) -> int:
    return sum(1 for r in rows if _is_number(r.get(field)))


def _detect_columns(fields: list[str], rows: list[dict[str, str]]) -> tuple[str, str, str]:
    by_norm = {_norm_header(f): f for f in fields}
    x_aliases: list[tuple[str, str]] = [
        ("cv", "CV"), ("columnvolume", "CV"), ("columnvolumes", "CV"),
        ("columnvolumecv", "CV"), ("elutionvolumecv", "CV"),
        ("columnvolumeml", "ML"), ("columnvolumesml", "ML"), ("columnvolumeinml", "ML"),
        ("elutionvolumeml", "ML"), ("retentionvolumeml", "ML"), ("volumeml", "ML"),
        ("elutionvolumeinml", "ML"), ("elutionvolume", "ML"), ("retentionvolume", "ML"), ("volume", "ML"), ("ml", "ML"),
        ("retentiontimemin", "MIN"), ("acquisitiontimemin", "MIN"), ("samplingtimemin", "MIN"), ("retentiontime", "MIN"), ("acquisitiontime", "MIN"), ("samplingtime", "MIN"), ("elapsedtimemin", "MIN"), ("elapsedtime", "MIN"),
        ("elutiontimemin", "MIN"), ("timemin", "MIN"), ("minutes", "MIN"), ("minute", "MIN"), ("min", "MIN"), ("time", "MIN"),
        ("retentiontimes", "S"), ("acquisitiontimes", "S"), ("samplingtimes", "S"), ("elapsedtimes", "S"), ("elutiontimes", "S"),
        ("timesec", "S"), ("timesecs", "S"), ("times", "S"), ("timeinseconds", "S"), ("seconds", "S"), ("second", "S"), ("sec", "S"),
        ("timeh", "H"), ("timehr", "H"), ("timehrs", "H"), ("timehours", "H"), ("hours", "H"),
        ("sample", "INDEX"), ("sampleid", "INDEX"), ("samplenumber", "INDEX"), ("sampleindex", "INDEX"),
        ("datapoint", "INDEX"), ("point", "INDEX"), ("frame", "INDEX"), ("sequence", "INDEX"),
        ("acquisitionnumber", "INDEX"), ("index", "INDEX"), ("scan", "INDEX"), ("scannumber", "INDEX"),
        ("x", "CV"),
    ]
    numeric_fields = [f for f in fields if _column_numeric_count(rows, f) >= 3]
    # A point-weight column is not a detector response or x-axis.
    numeric_fields = [f for f in numeric_fields if _norm_header(f) not in {"weight", "weights", "pointweight", "sampleweight", "lsqweight", "fitweight"}]
    # Prefer an explicitly labelled UV/mAU channel to generic "Signal" when
    # an instrument export contains several numerical detector channels.
    signal_aliases = ["uvmau", "uv280mau", "uv280", "uv", "absorbancemau"]
    signal_tags = ("uv", "a280", "mau", "absorbance")
    signal_field = ""
    for alias in signal_aliases:
        candidate = by_norm.get(alias)
        if candidate in numeric_fields:
            signal_field = candidate
            break
    if not signal_field:
        for field in fields:
            norm = _norm_header(field)
            if field in numeric_fields and any(tag in norm for tag in signal_tags):
                signal_field = field
                break
    if not signal_field:
        for alias in ("absorbance", "signal", "y", "totalproteingl", "totalgl", "proteingl", "concentrationgl", "response"):
            candidate = by_norm.get(alias)
            if candidate in numeric_fields:
                signal_field = candidate
                break
    if not signal_field:
        for field in fields:
            norm = _norm_header(field)
            if field in numeric_fields and any(tag in norm for tag in ("signal", "detector", "response", "intensity", "totalprotein", "concentrationgl", "yaxis")):
                signal_field = field
                break

    x_field = ""
    x_unit = "CV"
    for alias, unit in x_aliases:
        candidate = by_norm.get(alias)
        if candidate and candidate in numeric_fields and candidate != signal_field:
            x_field = candidate
            x_unit = unit
            break
    if not x_field:
        x_candidates = [f for f in numeric_fields if f != signal_field]
        if x_candidates:
            x_field = x_candidates[0]
            x_unit = "CV"
    if x_field and x_field != SAMPLE_ORDER_FIELD and x_unit == "MIN":
        if any(":" in str(row.get(x_field) or "") for row in rows):
            x_unit = "S"
    if not signal_field:
        candidates = [f for f in numeric_fields if f != x_field]
        if candidates:
            signal_field = candidates[0]
    if not signal_field:
        raise ValueError(
            "No numeric detector/signal column could be detected. Keep the detector values "
            "in a numeric CSV column (for example UV 280 [mAU] or Signal)."
        )
    if not x_field:
        # A single detector column is still usable when sampling was uniform. Its
        # row order is mapped across the selected recipe's full process duration.
        x_field = SAMPLE_ORDER_FIELD
        x_unit = "INDEX"
    return x_field, signal_field, x_unit


def _select_columns(fields: list[str], rows: list[dict[str, str]], *,
                    x_column: str | None = None, signal_column: str | None = None,
                    x_unit: str | None = None) -> tuple[str, str, str]:
    auto_x, auto_signal, auto_unit = _detect_columns(fields, rows)
    by_header = {_norm_header(f): f for f in fields}

    def choose(requested: str | None, inferred: str, role: str) -> str:
        if not str(requested or "").strip():
            return inferred
        selected = by_header.get(_norm_header(str(requested)))
        if selected is None:
            raise ValueError(f"Requested {role} column {requested!r} does not exist. Available: {fields}")
        if _column_numeric_count(rows, selected) < 3:
            raise ValueError(f"Requested {role} column {requested!r} does not contain at least three numeric samples.")
        return selected

    x = choose(x_column, auto_x, "x-axis")
    y = choose(signal_column, auto_signal, "signal")
    if x == y:
        raise ValueError("Reference x-axis and signal must be different columns.")
    unit = str(x_unit or "AUTO").strip().upper()
    if unit == "AUTO":
        if x != auto_x:
            raise ValueError("When overriding the automatically detected x-axis, also set the reference X-axis unit explicitly.")
        unit = auto_unit
    if unit not in {"CV", "ML", "MIN", "S", "H", "INDEX"}:
        raise ValueError("Reference x-axis unit must be AUTO, CV, ML, MIN, S, H, or INDEX.")
    return x, y, unit


def _conductivity_column(fields: list[str], rows: list[dict[str, str]],
                         requested: str | None, fit: bool) -> str | None:
    if not fit:
        return None
    if requested:
        matches = [field for field in fields if _norm_header(field) == _norm_header(requested)]
        if not matches:
            raise ValueError(f"Requested conductivity column {requested!r} does not exist. Available: {fields}")
        field = matches[0]
    else:
        matches = [field for field in fields if "conductivity" in _norm_header(field)
                   and _column_numeric_count(rows, field) >= 3]
        if len(matches) != 1:
            raise ValueError("Conductivity fitting requires one identifiable numeric conductivity column; enter its exact column heading for this run.")
        field = matches[0]
    if _column_numeric_count(rows, field) < 3:
        raise ValueError(f"Conductivity column {field!r} needs at least three numeric samples.")
    return field


def _conductivity_scale(field: str) -> float:
    name = field.casefold().replace("µ", "u").replace("μ", "u")
    if re.search(r"(?<![a-z])us\s*/\s*cm(?![a-z])", name):
        return 0.001
    if re.search(r"(?<![a-z])s\s*/\s*m(?![a-z])", name):
        return 10.0
    # Instrument exports commonly name an already-calibrated mS/cm trace
    # simply "Conductivity". Explicit column selection uses the same rule.
    return 1.0


def inspect_chromatogram_csv(path: Path | str, *, sheet_name: str | None = None,
                             x_column: str | None = None, signal_column: str | None = None,
                             x_unit: str | None = None, x_origin: str | None = None,
                             require_uv: bool = False, fit_conductivity: bool = False,
                             conductivity_column: str | None = None) -> dict[str, Any]:
    p = _resolve_chromatogram_path(path)
    fields, rows, actual_sheet = _read_numeric_rows(p, sheet_name=sheet_name)
    x_field, signal_field, unit = _select_columns(fields, rows, x_column=x_column, signal_column=signal_column, x_unit=x_unit)
    if require_uv:
        _uv_reference_scale(signal_field, signal_column)
    cond_field = _conductivity_column(fields, rows, conductivity_column, fit_conductivity)
    if cond_field == x_field or cond_field == signal_field:
        raise ValueError("The conductivity column must differ from the X-axis and UV columns.")
    origin = _validate_x_origin(x_origin, unit)
    paired = (
        _column_numeric_count(rows, signal_field)
        if x_field == SAMPLE_ORDER_FIELD else
        sum(1 for r in rows if _is_number(r.get(x_field)) and _is_number(r.get(signal_field)))
    )
    if paired < 3:
        raise ValueError("Detected chromatogram columns do not contain at least three paired numeric points.")
    assumption = None
    if x_field == SAMPLE_ORDER_FIELD:
        assumption = "No numeric time/CV/volume axis was present; sample order is assumed uniformly spaced across this profile's full process duration."
    return {
        "path": str(p),
        "rows": paired,
        "x_column": "sample order (assumed)" if x_field == SAMPLE_ORDER_FIELD else x_field,
        "signal_column": signal_field,
        "x_unit": unit,
        "x_origin": origin,
        "sheet_name": actual_sheet or "CSV",
        "columns": fields,
        "x_axis_assumption": assumption,
        "weight_column": _weight_column(fields),
        "conductivity_column": cond_field,
    }


def attach_chromatogram_csv(source: Path | str, **column_options: Any) -> dict[str, Any]:
    source_path = _resolve_chromatogram_path(source)
    info = inspect_chromatogram_csv(source_path, **column_options)
    # Profile paths are saved directly in their independent LSQ process table.
    # Never copy Excel bytes to a .csv filename or overwrite another run's data.
    info["source_path"] = str(source_path)
    return info


def _weight_column(fields: list[str]) -> str | None:
    return next((f for f in fields if _norm_header(f) in {
        "weight", "weights", "pointweight", "sampleweight", "lsqweight", "fitweight"}), None)


def _x_to_cv(x: np.ndarray, x_unit: str, config: dict[str, Any]) -> np.ndarray:
    unit = str(x_unit).upper()
    if unit == "CV":
        return x
    V = float(config["column"]["volume_mL"])
    if unit == "ML":
        return x / V
    if unit in {"MIN", "S", "H", "INDEX"}:
        # Flow may change between load, PLW and elution stages.  Convert time
        # using the exact same piecewise program as the mechanistic solver.
        program = _build_program(config, np.zeros(max(1, len(active_components(config))), dtype=float))
        values = np.asarray(x, dtype=float)
        if unit == "INDEX":
            span = float(np.max(values) - np.min(values))
            fraction = np.zeros_like(values) if span <= 0 else (values - float(np.min(values))) / span
            seconds = fraction * program.total_time_s
        else:
            seconds = values * (3600.0 if unit == "H" else 60.0 if unit == "MIN" else 1.0)
        return np.asarray([
            float(t) * program.stages[0].flow_mL_min / (60.0 * V) if t < 0.0 else
            program.total_CV + (float(t) - program.total_time_s) * program.stages[-1].flow_mL_min / (60.0 * V) if t > program.total_time_s else
            program.time_to_cv(float(t)) for t in seconds
        ], dtype=float)
    raise ValueError(f"Unsupported chromatogram x-axis unit: {x_unit}")


def _validate_x_origin(x_origin: str | None, x_unit: str) -> str:
    origin = str(x_origin or "RUN_START").strip().upper()
    if origin not in {"RUN_START", "ELUTION_START"}:
        raise ValueError("Reference X-axis zero must be RUN_START or ELUTION_START.")
    if origin == "ELUTION_START" and x_unit not in {"CV", "ML"}:
        raise ValueError("Elution-start X-axis zero requires CV or mL units.")
    return origin


def _elution_start_cv(config: dict[str, Any]) -> float:
    program = _build_program(config, np.zeros(max(1, len(active_components(config))), dtype=float))
    for index, stage in enumerate(program.stages):
        if stage.kind == "ELUTION":
            return float(program.stage_start_CV[index])
    raise ValueError("The selected LSQ process recipe has no elution stage; elution-start X-axis zero cannot be used.")


def load_chromatogram_reference(path: Path | str, config: dict[str, Any], *,
                                sheet_name: str | None = None, x_column: str | None = None,
                                signal_column: str | None = None, x_unit: str | None = None,
                                x_origin: str | None = None, require_uv: bool = False,
                                fit_conductivity: bool = False,
                                conductivity_column: str | None = None) -> ChromatogramReference:
    p = _resolve_chromatogram_path(path)
    fields, rows, actual_sheet = _read_numeric_rows(p, sheet_name=sheet_name)
    x_field, signal_field, unit = _select_columns(fields, rows, x_column=x_column, signal_column=signal_column, x_unit=x_unit)
    signal_scale = _uv_reference_scale(signal_field, signal_column) if require_uv else 1.0
    cond_field = _conductivity_column(fields, rows, conductivity_column, fit_conductivity)
    if cond_field == x_field or cond_field == signal_field:
        raise ValueError("The conductivity column must differ from the X-axis and UV columns.")
    origin = _validate_x_origin(x_origin, unit)
    weights_field = _weight_column(fields)
    pairs: list[tuple[float, float, float]] = []
    if x_field == SAMPLE_ORDER_FIELD:
        for row in rows:
            signal = _parse_numeric_cell(row.get(signal_field))
            if signal is not None:
                weight = _reference_row_weight(row, weights_field)
                pairs.append((float(len(pairs)), signal, weight))
    else:
        for row in rows:
            x_value = _parse_numeric_cell(row.get(x_field))
            signal = _parse_numeric_cell(row.get(signal_field))
            if x_value is not None and signal is not None:
                weight = _reference_row_weight(row, weights_field)
                pairs.append((x_value, signal, weight))
    if len(pairs) < 3:
        raise ValueError("Chromatogram CSV contains fewer than three paired numeric points.")
    arr = np.asarray(pairs, dtype=float)
    order = np.argsort(arr[:, 0])
    x = arr[order, 0]
    y = arr[order, 1]
    y = y * signal_scale
    weights = arr[order, 2]
    cv = _x_to_cv(x, unit, config)
    if origin == "ELUTION_START":
        cv = cv + _elution_start_cv(config)
    finite = np.isfinite(cv) & np.isfinite(y)
    cv, y, weights = cv[finite], y[finite], weights[finite]
    if len(cv) < 3:
        raise ValueError("Chromatogram CSV contains fewer than three finite points after conversion to CV.")
    # Repeated reference coordinates are retained: every measured sample is
    # one equation-9 residual, including replicates at the same time/CV.
    display_x_field = "sample order (assumed)" if x_field == SAMPLE_ORDER_FIELD else x_field
    assumption = None
    if x_field == SAMPLE_ORDER_FIELD:
        assumption = "No numeric time/CV/volume axis was present; sample order is assumed uniformly spaced across the selected profile's process duration."
    elif unit == "INDEX":
        assumption = "The numeric sample/point index is mapped uniformly across the selected profile's process duration."
    if not np.any(weights > 0):
        raise ValueError("All reference sample weights are zero; at least one must be positive.")
    cond_cv = cond_signal = None
    if cond_field:
        cond_pairs = []
        for row_index, row in enumerate(rows):
            raw_x = float(row_index) if x_field == SAMPLE_ORDER_FIELD else _parse_numeric_cell(row.get(x_field))
            raw_cond = _parse_numeric_cell(row.get(cond_field))
            if raw_x is not None and raw_cond is not None:
                cond_pairs.append((raw_x, raw_cond * _conductivity_scale(cond_field)))
        if len(cond_pairs) < 3:
            raise ValueError("Conductivity fitting needs at least three paired numeric X/conductivity points.")
        cond_arr = np.asarray(cond_pairs, dtype=float)
        cond_arr = cond_arr[np.argsort(cond_arr[:, 0], kind="stable")]
        if x_field == SAMPLE_ORDER_FIELD:
            # Preserve each row's position in the full instrument export even
            # when conductivity has gaps.
            cond_cv = cond_arr[:, 0] / max(1, len(rows) - 1) * _build_program(
                config, np.zeros(max(1, len(active_components(config))), dtype=float)).total_CV
        elif unit == "INDEX":
            # The conductivity channel may be sparse. Its index coordinates
            # must use the UV channel's full reference index span.
            span = float(x[-1] - x[0])
            fraction = np.zeros(len(cond_arr), dtype=float) if span <= 0 else (cond_arr[:, 0] - x[0]) / span
            cond_cv = fraction * _build_program(
                config, np.zeros(max(1, len(active_components(config))), dtype=float)).total_CV
        else:
            cond_cv = _x_to_cv(cond_arr[:, 0], unit, config)
        if origin == "ELUTION_START":
            cond_cv = cond_cv + _elution_start_cv(config)
        cond_signal = cond_arr[:, 1]
        if not np.all(np.isfinite(cond_cv) & np.isfinite(cond_signal)):
            raise ValueError("Conductivity reference contains non-finite converted values.")
    return ChromatogramReference(cv=cv, signal=y, x_column=display_x_field, signal_column=signal_field, x_unit=unit, source_path=str(p), x_axis_assumption=assumption, point_weights=weights, weight_column=weights_field, conductivity_cv=cond_cv, conductivity_signal=cond_signal, conductivity_column=cond_field, sheet_name=actual_sheet)


def _reference_row_weight(row: dict[str, str], weight_field: str | None) -> float:
    if weight_field is None:
        return 1.0
    weight = _parse_numeric_cell(row.get(weight_field))
    if weight is None or weight < 0 or not math.isfinite(weight):
        raise ValueError(f"Point-weight column {weight_field!r} must contain finite nonnegative values for every paired sample.")
    return weight


def _positive_log_spec(path: str, label: str, value: float, decades: float = 1.5) -> FitParameter:
    if value <= 0 or not math.isfinite(value):
        raise ValueError(f"{label} must be positive before least-squares refinement.")
    x = math.log10(value)
    return FitParameter(path, label, "LOG10", x - decades, x + decades, x)


def _is_concentration_signal(signal_column: str) -> bool:
    signal_norm = _norm_header(signal_column)
    return any(tag in signal_norm for tag in ("totalproteingl", "totalgl", "proteingl", "concentrationgl"))


def _uv_reference_scale(signal_column: str, selected_column: str | None) -> float:
    """Check that an LSQ reference is UV and return its conversion to mAU."""
    name = _norm_header(signal_column)
    non_uv = ("conductivity", "salt", "pressure", "flow", "temperature", "turbidity",
              "percentb", "masspercent", "masspct", "concentration", "totalprotein", "mgml")
    if (_is_concentration_signal(signal_column) or any(tag in name for tag in non_uv)
            or name in {"ph", "weight", "pointweight", "sampleweight"}):
        raise ValueError(f"Least-squares refinement requires a UV signal, not {signal_column!r}.")
    if "voltage" in name or re.search(r"(?<![a-z])m?v(?![a-z])", signal_column.casefold()):
        raise ValueError(f"UV reference {signal_column!r} must use absorbance units (AU or mAU), not voltage.")
    is_labelled_uv = any(tag in name for tag in ("uv", "a280", "mau", "absorbance"))
    if not is_labelled_uv and not str(selected_column or "").strip():
        raise ValueError(
            f"Cannot identify {signal_column!r} as UV. Choose the UV signal column explicitly "
            "or label it UV/mAU in the file. Generic selected signals are interpreted as mAU."
        )
    # Absorbance exports may use AU while the model reports mAU.
    au_label = re.search(r"(?<![a-z])au(?![a-z])", signal_column.casefold())
    return 1000.0 if au_label or (name.endswith("au") and not name.endswith("mau")) else 1.0


def _linear_z_spec(path: str, label: str, value: float) -> FitParameter:
    span = max(25.0, abs(value) * 1.5)
    return FitParameter(path, label, "LINEAR", value - span, value + span, value)


def _default_unlocked_parameter_paths(config: dict[str, Any], *, fit_uv_response: bool = True) -> set[str]:
    """Start TDM–CPA with target retention parameters; honor saved lock states."""
    config = normalize_config(config)
    model = str(config["model"])
    if model == MODEL_TDM_CPA:
        return {"components[0].delta_ref", "components[0].Z_ref"}
    paths: set[str] = set()
    for i, row in enumerate(selected_components(config)):
        if float(row.get("mass_percent") or 0) <= 0:
            continue
        prefix = f"components[{i}]"
        if model == MODEL_TDM_WANG:
            paths.update({f"{prefix}.wang_k_eq", f"{prefix}.wang_beta1_per_M"})
        elif model in {MODEL_EDM_LANGMUIR, MODEL_TDM_LANGMUIR}:
            paths.update({f"{prefix}.qmax_g_L", f"{prefix}.b_L_g", f"{prefix}.salt_sensitivity_per_M"})
        else:
            paths.update({f"{prefix}.delta_ref", f"{prefix}.As_m_inv"})
            paths.add(f"{prefix}.salt_sensitivity_per_M" if is_hic(config) else f"{prefix}.Z_ref")
            if model == MODEL_TDM_CPA:
                paths.add(f"{prefix}.kkin_star_s")
        if i > 0 and fit_uv_response:
            paths.add(f"{prefix}.uv_response_factor_mAU_L_g")
    return paths


def _all_fit_parameter_specs(config: dict[str, Any], *, fit_uv_response: bool = True,
                             fit_conductivity: bool = False,
                             candidate_paths: Iterable[str] | None = None) -> list[FitParameter]:
    """Build bounded specs for every required mechanistic input except column geometry.

    Feed composition, load/flow, buffers, PLWs and elution are experimental
    inputs and remain fixed per profile. Column volume and length are also fixed
    and are never eligible for least-squares fitting.
    """
    config = normalize_config(config)
    model = str(config["model"])
    required_paths = required_parameter_paths(config)
    active_paths = set(required_paths)
    paths = [canonical_fit_parameter_path(p) for p in candidate_paths] if candidate_paths is not None else required_paths
    # A caller may supply a broader catalog (or saved locks from a previous
    # model). Never fit a parameter absent from the selected model's equations.
    paths = [path for path in paths if path in active_paths]
    specs: list[FitParameter] = []
    labels = {
        "buffer_dispersion_mL": "Buffer mixer dispersion volume [mL]",
        "salt_dispersion_mm2_s": "Buffer/salt axial dispersion [mm²/s]",
        "uv_dead_volume_mL": "UV system dead volume [mL]",
        "conductivity_dead_volume_mL": "Conductivity system dead volume [mL]",
        "uv_to_protein_mAU_L_g": "Target UV → protein response factor [mAU·L/g]",
        "conductivity_to_salt_M_per_mS_cm": "Conductivity → salt conversion factor [M per (mS/cm)]",
        "bead_radius_um": "Bead radius r_p [µm]",
        "void_fraction": "Void fraction ε_v [-]",
        "particle_porosity": "Particle porosity ε_p [-]",
        "salt_axial_dispersion_mm2_s": "Salt axial dispersion D_ax,salt [mm²/s]",
        "total_bed_porosity": "Total bed porosity ε [-]",
        "ligand_surface_density_umol_m2": "Ligand surface density Γ_L [µmol/m²]",
        "system_specific_adsorption_parameter": "System-specific adsorption parameter [-]",
        "D_ax_mm2_s": "Axial dispersion D_ax,i [mm²/s]",
        "k_eff_um_s": "Effective mass transfer k_eff,i [µm/s]",
        "accessible_particle_porosity": "Accessible particle porosity ε_p,i [-]",
        "D_app_mm2_s": "Apparent dispersion D_app,i [mm²/s]",
        "qmax_g_L": "Saturation capacity qmax,i [g/L stationary phase]",
        "b_L_g": "Competitive Langmuir b_i [L/g]",
        "salt_sensitivity_per_M": "HIC salt sensitivity k_s,i [M⁻¹]",
        "diameter_nm": "Protein diameter [nm]",
        "As_m_inv": "Accessible adsorber surface A_s,i [m⁻¹]",
        "Z_ref": "Protein charge Z_i [-]",
        "delta_ref": "CPA interaction parameter Δ_i [-]",
        "kkin_star_s": "CPA kinetic parameter k*kin,i [s⁻¹]",
        "pH_ref": "CPA reference pH",
        "Z1_per_pH": "CPA Z1,i [-]",
        "Z2_per_pH2": "CPA Z2,i [-]",
        "Z3_per_pH3": "CPA Z3,i [-]",
        "delta_pH_slope_m2_C": "CPA Δ1,i [m²/C]",
        "uv_response_factor_mAU_L_g": "UV → protein response factor [mAU·L/g]",
        "q0_g_L": "Wang q0 [g/L stationary phase]",
        "eta": "Wang η [-]",
        "wang_K_kin_s": "Wang K′kin,i [s]",
        "wang_k_eq": "Wang k_eq,i",
        "wang_qmax_g_L": "Wang qmax,i [g/L stationary phase]",
        "wang_n": "Wang n_i [-]",
        "wang_beta0": "Wang β0,i [-]",
        "wang_beta1_per_M": "Wang β1,i [M⁻¹]",
        "wang_beta2_L_g": "Wang β2,i [L/g]",
        "wang_beta3_per_pH": "Wang β3,i [pH⁻¹]",
    }
    for path in paths:
        if path in {"column.volume_mL", "column.length_mm"} or path.endswith(".mass_percent"):
            continue
        if path == "system.conductivity_dead_volume_mL" and not fit_conductivity:
            continue
        if not fit_uv_response and (path.endswith("uv_response_factor_mAU_L_g") or path == "conversion.uv_to_protein_mAU_L_g"):
            continue
        try:
            value = float(_get_path(config, path))
        except (KeyError, TypeError, ValueError):
            continue
        if not math.isfinite(value):
            continue
        field = path.rsplit(".", 1)[-1]
        if field in {"void_fraction", "particle_porosity", "total_bed_porosity", "accessible_particle_porosity"}:
            specs.append(FitParameter(path, labels.get(field, field), "LINEAR", 1e-9, 0.999999,
                                      float(np.clip(value, 1e-9, 0.999999))))
        elif field == "salt_sensitivity_per_M":
            specs.append(FitParameter(path, labels[field], "LINEAR", 0.0 if is_hic(config) else -25.0, 25.0,
                                      float(np.clip(value, 0.0 if is_hic(config) else -25.0, 25.0))))
        elif field in {"eta", "wang_n"}:
            lower, upper = (0.1, 3.0) if field == "eta" else (0.1, 4.0)
            specs.append(FitParameter(path, labels[field], "LINEAR", lower, upper,
                                      float(np.clip(value, lower, upper))))
        elif field in {"wang_beta0", "wang_beta1_per_M", "wang_beta2_L_g", "wang_beta3_per_pH"}:
            lower, upper = {"wang_beta0": (0.0, 2.0), "wang_beta1_per_M": (-5.0, 5.0),
                            "wang_beta2_L_g": (-5.0, 5.0), "wang_beta3_per_pH": (-1.0, 1.0)}[field]
            specs.append(FitParameter(path, labels[field], "LINEAR", lower, upper,
                                      float(np.clip(value, lower, upper))))
        elif field in {"buffer_dispersion_mL", "salt_dispersion_mm2_s", "uv_dead_volume_mL", "conductivity_dead_volume_mL", "D_ax_mm2_s", "D_app_mm2_s", "salt_axial_dispersion_mm2_s"}:
            upper = max(10.0, value * 20.0, float(config["column"]["volume_mL"]) * 5.0)
            specs.append(FitParameter(path, labels.get(field, field), "LINEAR", 0.0, upper,
                                      value))
        elif field == "Z_ref":
            specs.append(FitParameter(path, labels[field], "LINEAR", -500.0, 500.0,
                                      float(np.clip(value, -500.0, 500.0))))
        elif field == "pH_ref":
            specs.append(FitParameter(path, labels[field], "LINEAR", 0.0, 14.0,
                                      float(np.clip(value, 0.0, 14.0))))
        elif field in {"Z1_per_pH", "Z2_per_pH2", "Z3_per_pH3", "delta_pH_slope_m2_C"}:
            span = max(1.0, abs(value) * 2.0)
            specs.append(FitParameter(path, labels[field], "LINEAR", value - span, value + span, value))
        elif field in {"uv_to_protein_mAU_L_g", "uv_response_factor_mAU_L_g", "conductivity_to_salt_M_per_mS_cm",
                       "bead_radius_um", "salt_axial_dispersion_mm2_s", "ligand_surface_density_umol_m2",
                       "system_specific_adsorption_parameter", "D_ax_mm2_s", "k_eff_um_s", "D_app_mm2_s",
                       "qmax_g_L", "b_L_g", "diameter_nm", "As_m_inv", "delta_ref", "kkin_star_s",
                       "q0_g_L", "wang_K_kin_s", "wang_k_eq", "wang_qmax_g_L"}:
            try:
                specs.append(_positive_log_spec(path, labels.get(field, field), value))
            except ValueError:
                continue
    return [FitParameter(spec.path, _report_parameter_label(config, spec.path, spec.label),
                         spec.transform, spec.lower, spec.upper, spec.initial) for spec in specs]


def _report_parameter_label(config: dict[str, Any], path: str, field_label: str | None = None) -> str:
    match = re.fullmatch(r"components\[(\d+)\]\.([A-Za-z0-9_]+)", path)
    if not match:
        return field_label or path
    index = int(match.group(1))
    # Parameter paths address the saved component slots, whereas the solver's
    # active list may omit a zero-mass impurity between populated slots.
    species = config.get("components", [])
    if index >= len(species):
        return field_label or path
    role = "Target protein" if index == 0 else f"Impurity {index}"
    name = str(species[index].get("name") or role)
    return f"{role} ({name}) — {field_label or match.group(2)}"


def _limit_fit_specs(config: dict[str, Any], specs: list[FitParameter]) -> list[FitParameter]:
    """Intersect optional ±% movement limits with the model's physical bounds."""
    limits = config.get("least_squares", {}).get("fit_range_percent") or {}
    if not isinstance(limits, dict):
        raise ValueError("Per-parameter LSQ fit ranges must be a path-to-percentage mapping.")
    limited: list[FitParameter] = []
    for spec in specs:
        entered = limits.get(spec.path)
        if entered is None:
            limited.append(spec)
            continue
        try:
            percent = float(entered)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{spec.label}: ±% fit range must be numeric.") from exc
        if not math.isfinite(percent) or not 0 < percent <= 1000:
            raise ValueError(f"{spec.label}: ±% fit range must be greater than 0 and at most 1000.")
        fraction = percent / 100.0
        if spec.transform == "LOG10":
            start = 10.0 ** spec.initial
            requested_lower = (math.log10(start * (1.0 - fraction))
                               if fraction < 1.0 else spec.lower)
            requested_upper = math.log10(start * (1.0 + fraction))
        else:
            # A percentage of zero would freeze a parameter. Use the physical
            # default bound span for a zero-valued starting parameter.
            delta = fraction * (abs(spec.initial) if spec.initial != 0
                                else spec.upper - spec.lower)
            requested_lower = spec.initial - delta
            requested_upper = spec.initial + delta
        lower = max(spec.lower, requested_lower)
        upper = min(spec.upper, requested_upper)
        if not lower < upper:
            raise ValueError(f"{spec.label}: ±% fit range leaves no feasible movement within the physical bounds.")
        limited.append(FitParameter(spec.path, spec.label, spec.transform,
                                    lower, upper, spec.initial))
    return limited


def _physical_fit_bounds(spec: FitParameter) -> tuple[float, float]:
    if spec.transform == "LOG10":
        return 10.0 ** spec.lower, 10.0 ** spec.upper
    return spec.lower, spec.upper


def fit_parameter_lock_catalog(config: dict[str, Any]) -> list[FitParameter]:
    """Return all currently required mechanistic parameters that can be unlocked."""
    return _all_fit_parameter_specs(config, fit_uv_response=True, fit_conductivity=True)


def fit_parameter_specs(config: dict[str, Any], *, fit_uv_response: bool = True,
                        fit_conductivity: bool = False,
                        unlocked_paths: Iterable[str] | None = None,
                        candidate_paths: Iterable[str] | None = None) -> list[FitParameter]:
    """Return the model-backed parameters refined by least squares.

    Eligible instrument, adsorption and column-transport values are refined only
    when unlocked. Recipe inputs, geometry and species composition remain fixed.
    """
    config = normalize_config(config)
    model = str(config["model"])
    if unlocked_paths is not None:
        unlocked = set(canonical_fit_parameter_path(path) for path in unlocked_paths)
        candidates = _all_fit_parameter_specs(config, fit_uv_response=fit_uv_response,
                                              fit_conductivity=fit_conductivity,
                                              candidate_paths=candidate_paths)
        return _limit_fit_specs(config, [spec for spec in candidates if spec.path in unlocked])
    persisted = config.get("least_squares", {}).get("unlocked_paths")
    unlocked = set(persisted) if persisted is not None else _default_unlocked_parameter_paths(config, fit_uv_response=fit_uv_response)
    return _limit_fit_specs(config, [
        spec for spec in _all_fit_parameter_specs(
            config, fit_uv_response=fit_uv_response,
            fit_conductivity=fit_conductivity, candidate_paths=candidate_paths)
        if spec.path in unlocked
    ])


def fit_parameter_text(config: dict[str, Any]) -> str:
    specs = fit_parameter_specs(config)
    return "\n".join(f"• {s.label}" for s in specs)


def _apply_vector(base: dict[str, Any], specs: list[FitParameter], vector: np.ndarray) -> dict[str, Any]:
    cfg = copy.deepcopy(base)
    for spec, raw in zip(specs, np.asarray(vector, dtype=float)):
        value = 10.0 ** float(raw) if spec.transform == "LOG10" else float(raw)
        _set_path(cfg, spec.path, value)
    cfg = normalize_config(cfg)
    for spec, raw in zip(specs, np.asarray(vector, dtype=float)):
        expected = 10.0 ** float(raw) if spec.transform == "LOG10" else float(raw)
        actual = float(_get_path(cfg, spec.path))
        if not math.isclose(actual, expected, rel_tol=1e-10, abs_tol=1e-12):
            raise ValueError(f"Parameter transfer failed for {spec.path}: requested {expected}, read back {actual}.")
    return cfg


def _encode_initial(specs: list[FitParameter]) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    x0 = np.array([s.initial for s in specs], dtype=float)
    lo = np.array([s.lower for s in specs], dtype=float)
    hi = np.array([s.upper for s in specs], dtype=float)
    return x0, lo, hi


def validate_composition_references(
    config: dict[str, Any],
    references: Iterable[dict[str, Any]],
) -> list[dict[str, Any]]:
    names = [str(r["name"]) for r in active_components(config)]
    if len(names) != len(set(names)):
        raise ValueError("Active species names must be unique for fraction mass% refinement.")
    out: list[dict[str, Any]] = []
    previous_cv = -math.inf
    previous_end_cv = -math.inf
    previous_index = 0
    for k, row in enumerate(references, start=1):
        try:
            fraction_index = int(row.get("fraction_index", k))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Fraction {k}: row number must be a positive integer.") from exc
        if fraction_index <= previous_index:
            raise ValueError(f"Fraction {k}: row numbers must increase.")
        previous_index = fraction_index
        mass_values = row.get("mass_percent") or {}
        if not isinstance(mass_values, dict):
            raise ValueError(f"Fraction {k}: species mass% values must be named columns.")
        unknown_names = set(mass_values) - set(names)
        if unknown_names:
            raise ValueError(f"Fraction {k}: unknown species {', '.join(sorted(unknown_names))}.")
        unit = str(row.get("x_unit") or "CV").strip().upper()
        if unit not in {"CV", "ML"}:
            raise ValueError(f"Fraction {k}: axis unit must be CV or mL.")
        legacy_point = "CV" in row and "start" not in row and "end" not in row
        if legacy_point:
            try:
                cv = float(row["CV"])
            except (TypeError, ValueError, KeyError) as exc:
                raise ValueError(f"Composition reference {k}: CV is required and must be numeric.") from exc
            start = end = cv
        else:
            try:
                start = float(row["start"])
                end = float(row["end"])
            except (TypeError, ValueError, KeyError) as exc:
                raise ValueError(f"Fraction {k}: enter both start and end in {unit}.") from exc
            if not math.isfinite(start) or not math.isfinite(end) or start < 0 or end <= start:
                raise ValueError(f"Fraction {k}: start must be >= 0 and end must be greater than start.")
            column_volume = float(config["column"]["volume_mL"])
            cv = (start + end) / (2.0 * column_volume) if unit == "ML" else (start + end) / 2.0
        if not math.isfinite(cv) or cv < 0:
            raise ValueError(f"Fraction {k}: midpoint CV must be finite and >= 0.")
        if cv <= previous_cv:
            raise ValueError(f"Fraction {k}: fraction midpoints must increase in row order.")
        previous_cv = cv
        start_cv = start / float(config["column"]["volume_mL"]) if unit == "ML" else start
        end_cv = end / float(config["column"]["volume_mL"]) if unit == "ML" else end
        if start_cv < previous_end_cv - 1e-9:
            raise ValueError(f"Fraction {k}: fraction intervals must be ordered and must not overlap.")
        previous_end_cv = end_cv
        measured: dict[str, float] = {}
        for name in names:
            raw = mass_values.get(name)
            if raw is None or str(raw).strip() == "":
                continue
            try:
                value = float(raw)
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Fraction {k}: mass% for {name} must be numeric or blank.") from exc
            if not math.isfinite(value) or value < 0 or value > 100:
                raise ValueError(f"Fraction {k}: mass% for {name} must be between 0 and 100.")
            measured[name] = value
        # Blank cells intentionally mean unmeasured species, so a partial row
        # need not total 100%. Even a partial row cannot exceed the available
        # 100% mass-fraction budget, allowing modest assay rounding.
        if sum(measured.values()) > 100.5:
            raise ValueError(f"Fraction {k}: entered species mass% total exceeds 100%. Check the assay basis and species mapping.")
        out.append({"CV": cv, "fraction_index": fraction_index,
                    "start": start, "end": end, "x_unit": unit,
                    "legacy_point": legacy_point,
                    "mass_percent": measured})
    if not any(ref["mass_percent"] for ref in out):
        raise ValueError("Enter at least one target or impurity mass% observation for the chromatogram + mass% option.")
    return out


def _baseline_correct(signal: np.ndarray) -> tuple[np.ndarray, float, float]:
    y = np.asarray(signal, dtype=float)
    n0 = max(3, min(len(y), int(math.ceil(len(y) * 0.05))))
    baseline = float(np.median(y[:n0]))
    corrected = y - baseline
    scale = max(float(np.max(corrected) - np.min(corrected)), float(np.max(np.abs(corrected))), 1e-12)
    return corrected, baseline, scale


def _chromatogram_residuals(config: dict[str, Any], ref: ChromatogramReference, trace: dict[str, Any]) -> tuple[np.ndarray, dict[str, Any]]:
    _uv_reference_scale(ref.signal_column, ref.signal_column)
    sim_cv = np.asarray(trace["CV"], dtype=float)
    sim_total = np.asarray(trace["total_g_L"], dtype=float)
    tolerance = 1e-9 * max(1.0, float(np.max(sim_cv)))
    mask = (ref.cv >= float(np.min(sim_cv)) - tolerance) & (ref.cv <= float(np.max(sim_cv)) + tolerance)
    if not np.all(mask):
        raise ValueError("Reference chromatogram extends outside the simulated recipe. Extend the run/flush duration or crop the reference explicitly; fitting will not silently discard measured points.")
    if int(np.count_nonzero(mask)) < 3:
        raise ValueError("Fewer than three chromatogram points overlap the simulated CV range.")
    x = np.clip(ref.cv[mask], float(np.min(sim_cv)), float(np.max(sim_cv)))
    y_meas = ref.signal[mask]
    y_sim = np.interp(x, sim_cv, sim_total)
    _meas_corr, estimated_baseline, measured_scale = _baseline_correct(y_meas)
    options = config.get("least_squares", {})
    baseline = estimated_baseline if options.get("baseline_mode") == "INITIAL_MEDIAN" else float(options.get("detector_baseline", 0.0))
    objective = options.get("objective", "RAW_SSE")
    normalized = objective == "NORMALIZED_MSE"
    scale = measured_scale if normalized else 1.0
    detector_scale = None
    predicted_zero_baseline = np.interp(x, sim_cv, np.asarray(trace["uv_mAU"], dtype=float))
    detector_scale_source = "species-specific UV response factors (target factor plus fitted impurity factors)"
    predicted_signal = baseline + predicted_zero_baseline
    residual = (predicted_signal - y_meas) / scale
    if normalized:
        residual = residual / math.sqrt(len(residual))
    if objective == "WEIGHTED_RMSE":
        point_weights = ref.point_weights if ref.point_weights is not None else np.ones(len(residual))
        residual = residual * np.sqrt(point_weights / float(np.sum(point_weights)))
    info = {
        "CV": x,
        "measured": y_meas,
        "simulated_total_g_L": y_sim,
        "predicted_detector_signal": predicted_signal,
        "baseline": baseline,
        "detector_scale": detector_scale,
        "detector_scale_source": detector_scale_source,
        "normalization_scale": scale,
        "measured_signal_scale": measured_scale,
        "objective": options.get("objective", "RAW_SSE"),
        "overlap_points": len(x),
        "point_weights": ref.point_weights if ref.point_weights is not None else np.ones(len(x)),
        "point_weight_column": ref.weight_column,
    }
    return residual, info


def _conductivity_residuals(config: dict[str, Any], ref: ChromatogramReference,
                            trace: dict[str, Any], uv_info: dict[str, Any]) -> tuple[np.ndarray, dict[str, Any]]:
    if ref.conductivity_cv is None or ref.conductivity_signal is None:
        raise ValueError("Conductivity fitting is enabled, but this run has no measured conductivity values.")
    sim_cv = np.asarray(trace["CV"], dtype=float)
    measured_cv = np.asarray(ref.conductivity_cv, dtype=float)
    measured = np.asarray(ref.conductivity_signal, dtype=float)
    tolerance = 1e-9 * max(1.0, float(sim_cv[-1]))
    if np.any(measured_cv < sim_cv[0] - tolerance) or np.any(measured_cv > sim_cv[-1] + tolerance):
        raise ValueError("Measured conductivity extends outside this run's simulated process recipe.")
    predicted = np.interp(np.clip(measured_cv, sim_cv[0], sim_cv[-1]), sim_cv,
                          np.asarray(trace["conductivity_mS_cm"], dtype=float))
    difference = predicted - measured
    span = float(np.max(measured) - np.min(measured))
    scale = span if span > 1e-6 else max(float(np.max(np.abs(measured))), 1e-3)
    objective = config["least_squares"]["objective"]
    residual = difference / (scale * math.sqrt(len(difference)))
    if objective == "RAW_SSE":
        # One conductivity channel has approximately the same influence as
        # one UV channel regardless of sampling rate or detector units.
        residual *= max(1.0, float(uv_info["measured_signal_scale"])) * math.sqrt(len(uv_info["CV"]))
    elif objective == "WEIGHTED_RMSE":
        residual *= max(1.0, float(uv_info["measured_signal_scale"]))
    return residual, {"CV": measured_cv, "measured": measured, "predicted": predicted,
                      "scale_mS_cm": scale, "column": ref.conductivity_column,
                      "rmse_mS_cm": float(np.sqrt(np.mean(difference ** 2)))}


def _reference_display_series(cv: np.ndarray, signal: np.ndarray, total_cv: float,
                              pixel_columns: int = 744) -> tuple[np.ndarray, np.ndarray]:
    """Show one robust reference value per horizontal pixel without changing fit data.

    Drawing thousands of raw samples into a few SVG pixels repeatedly paints
    nearly vertical black strokes, which can look like a filled peak. The full
    unbinned reference remains in the least-squares residual and trace CSV.
    """
    x = np.asarray(cv, dtype=float)
    y = np.asarray(signal, dtype=float)
    if len(x) != len(y):
        raise ValueError("Reference plot coordinates and UV values must be paired.")
    if not len(x):
        return x, y
    if not math.isfinite(total_cv) or total_cv <= 0:
        raise ValueError("The plotted run must have a positive finite CV span.")
    finite = np.isfinite(x) & np.isfinite(y)
    x, y = x[finite], y[finite]
    if not len(x):
        return x, y
    order = np.argsort(x, kind="stable")
    x, y = x[order], y[order]
    columns = np.clip(np.floor(x / total_cv * pixel_columns).astype(int), 0, pixel_columns - 1)
    boundaries = np.r_[0, np.flatnonzero(np.diff(columns)) + 1, len(columns)]
    return (
        np.array([np.median(x[a:b]) for a, b in zip(boundaries[:-1], boundaries[1:])]),
        np.array([np.median(y[a:b]) for a, b in zip(boundaries[:-1], boundaries[1:])]),
    )


def _chromatogram_fit_svg(config: dict[str, Any], info: dict[str, Any]) -> str:
    """Show the entire programmed run without extending measured UV samples."""
    chrom = info["chromatogram"]
    trace = info["result"]["trace"]
    program = info["result"]["simulation"]["program"]
    display_cv, display_signal = _reference_display_series(
        chrom["CV"], chrom["measured"], float(program.total_CV))
    return _svg_plot(
        trace["CV"],
        [("Measured UV", display_signal),
         ("Predicted UV", np.asarray(trace["uv_mAU"]) + chrom["baseline"])],
        series_x=[display_cv, trace["CV"]],
        stages=_stage_receipt(config, program), total_cv=program.total_CV,
        percent_B=trace["programmed_percent_B"],
        column_percent_B=trace["column_percent_B"],
        conductivity=trace["conductivity_mS_cm"],
        time_min=trace["time_min"], flow_mL_min=trace["flow_mL_min"],
        show_percent_B=config["system"]["show_percent_B"],
        show_conductivity=config["system"]["show_conductivity"],
        show_flow_rate=config["system"].get("show_flow_rate", True),
    )


def _conductivity_fit_svg(info: dict[str, Any]) -> str:
    cond = info["conductivity"]
    if not cond:
        return ""
    x = np.asarray(cond["CV"], dtype=float)
    measured = np.asarray(cond["measured"], dtype=float)
    predicted = np.asarray(cond["predicted"], dtype=float)
    x_max = max(float(info["result"]["simulation"]["program"].total_CV), 1e-9)
    low = float(min(np.min(measured), np.min(predicted)))
    high = float(max(np.max(measured), np.max(predicted)))
    pad = max((high - low) * 0.08, 0.01)
    low -= pad
    high += pad
    def xy(xs: np.ndarray, ys: np.ndarray) -> str:
        return " ".join(f"{56 + 770 * float(a) / x_max:.2f},{310 - 260 * (float(b) - low) / (high - low):.2f}"
                        for a, b in zip(xs, ys))
    disp_x, disp_y = _reference_display_series(x, measured, x_max)
    return (f'<h4>Measured vs predicted conductivity — {html.escape(str(cond["column"]))}</h4>'
            f'<p>Conductivity RMSE: {cond["rmse_mS_cm"]:.6g} mS/cm; '
            f'{len(x)} measured samples used in refinement.</p>'
            '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 900 365" '
            'role="img" aria-label="Measured and predicted conductivity by CV">'
            '<rect width="900" height="365" fill="white"/>'
            '<path d="M56 50V310H826" fill="none" stroke="#444"/>'
            f'<polyline points="{xy(disp_x, disp_y)}" fill="none" stroke="#111" stroke-width="1.8"/>'
            f'<polyline points="{xy(x, predicted)}" fill="none" stroke="#0878c6" stroke-width="1.8"/>'
            '<text x="56" y="338" font-size="12">0 CV</text>'
            f'<text x="826" y="338" text-anchor="end" font-size="12">{x_max:.4g} CV</text>'
            '<text x="60" y="28" font-size="12" fill="#111">■ Measured</text>'
            '<text x="180" y="28" font-size="12" fill="#0878c6">■ Predicted [mS/cm]</text></svg>')


def preview_lsq_alignment(
    profiles: Iterable[dict[str, Any]], *, output_html: Path | str = LSQ_ALIGNMENT_HTML,
) -> dict[str, Any]:
    """Overlay measured UV and the current model for each selected LSQ recipe.

    This uses the same reference reader, UV residual calculation and plot as
    refinement, but performs one simulation per run and no optimization.
    """
    profiles = list(profiles)
    if not 1 <= len(profiles) <= MAX_LSQ_PROFILES:
        raise ValueError(f"Alignment preview requires 1–{MAX_LSQ_PROFILES} selected LSQ runs.")
    sections: list[tuple[str, str]] = []
    metrics: list[dict[str, Any]] = []
    for index, profile in enumerate(profiles, start=1):
        config = normalize_config(profile["config"])
        context = dict(profile.get("run_context") or {})
        run_number = context.get("run_number", index)
        run_name = str(context.get("run_name") or f"LSQ Run {run_number}")
        errors = validate_config(config, strict=True)
        if errors:
            raise ValueError(f"LSQ Run {run_number} process recipe must be complete:\n- " + "\n- ".join(errors))
        source = profile.get("chromatogram_csv") or profile.get("csv_path")
        if not source:
            raise ValueError(f"LSQ Run {run_number} has no attached UV chromatogram.")
        reference = load_chromatogram_reference(
            source, config, sheet_name=profile.get("sheet_name"),
            x_column=profile.get("x_column"), signal_column=profile.get("signal_column"),
            x_unit=profile.get("x_unit"), x_origin=profile.get("x_origin"), require_uv=True,
            fit_conductivity=bool(profile.get("fit_conductivity")),
            conductivity_column=profile.get("conductivity_column"),
        )
        raw_comp_refs = profile.get("composition_references") or []
        comp_refs = validate_composition_references(config, raw_comp_refs) if raw_comp_refs else []
        _residual, info = _objective_parts(config, reference, MODE_CHROMATOGRAM, [], bool(profile.get("fit_conductivity")))
        if comp_refs:
            _composition_residuals(config, comp_refs, info["result"]["trace"])
        chrom = info["chromatogram"]
        predicted = np.asarray(chrom["predicted_detector_signal"], dtype=float)
        measured = np.asarray(chrom["measured"], dtype=float)
        rmse = float(np.sqrt(np.mean((predicted - measured) ** 2)))
        program = info["result"]["simulation"]["program"]
        program_end = float(program.total_CV)
        stage_receipt = "; ".join(
            f"{stage.name}: {program.stage_start_CV[i]:.6g}–{program.stage_end_CV[i]:.6g} CV"
            for i, stage in enumerate(program.stages)
        )
        initial = program.initial_chemistry or {}
        initial_receipt = (
            f"{initial.get('source', 'MATCH_LOAD')}: pH {float(initial.get('pH', 0)):.6g}, "
            f"salt {float(initial.get('salt_concentration_M', 0)):.6g} M, "
            f"ionic strength {float(initial.get('ionic_strength_M', 0)):.6g} M, "
            f"conductivity {float(initial.get('conductivity_mS_cm', 0)):.6g} mS/cm"
        )
        metric = {
            "run_number": run_number,
            "run_name": run_name,
            "model_run_label": str(context.get("model_run_label") or f"Independent LSQ recipe Run {run_number}"),
            "model_run_name": str(context.get("model_run_name") or run_name),
            "sample_points": int(len(measured)),
            "measured_start_CV": float(reference.cv[0]),
            "measured_end_CV": float(reference.cv[-1]),
            "programmed_end_CV": program_end,
            "uv_rmse_mAU": rmse,
            "conductivity_rmse_mS_cm": info["conductivity"]["rmse_mS_cm"] if info["conductivity"] else None,
            "mass_percent_observations": sum(len(row["mass_percent"]) for row in comp_refs),
            "x_column": reference.x_column,
            "x_unit": reference.x_unit,
            "x_axis_assumption": reference.x_axis_assumption or "RUN_START",
            "signal_column": reference.signal_column,
            "source_path": str(Path(source).resolve()),
            "sheet_name": reference.sheet_name or "CSV",
            "column_volume_mL": float(program.column_volume_mL),
            "stage_receipt": stage_receipt,
            "initial_equilibration": dict(initial),
        }
        metrics.append(metric)
        sections.append((str(run_number),
            f"<h2>LSQ Run {html.escape(str(run_number))}: {html.escape(run_name)}</h2>"
            f"<p><b>Model UV recipe:</b> {html.escape(metric['model_run_label'])} — "
            f"{html.escape(metric['model_run_name'])}. The model was simulated from the saved recipe for this preview.</p>"
            f"<p><b>Process setup:</b> column volume {metric['column_volume_mL']:.6g} mL; "
            f"{html.escape(stage_receipt)}.<br><b>Initial column equilibration:</b> "
            f"{html.escape(initial_receipt)}.</p>"
            f"<p><b>Attached Excel / CSV:</b> {html.escape(metric['source_path'])}; "
            f"<b>worksheet:</b> {html.escape(metric['sheet_name'])}<br>"
            f"<b>Axis:</b> {html.escape(reference.x_column)} [{html.escape(reference.x_unit)}], "
            f"{html.escape(metric['x_axis_assumption'])}; "
            f"<b>signal:</b> {html.escape(reference.signal_column)}<br>"
            f"<b>Measured span:</b> {metric['measured_start_CV']:.6g}–{metric['measured_end_CV']:.6g} CV "
            f"({metric['sample_points']} samples); <b>programmed end:</b> {program_end:.6g} CV; "
            f"<b>UV RMSE before refinement:</b> {rmse:.6g} mAU.</p>"
            + (f'<p><b>Measured conductivity:</b> {html.escape(str(reference.conductivity_column))}; '
               f'<b>RMSE before refinement:</b> {info["conductivity"]["rmse_mS_cm"]:.6g} mS/cm.</p>'
               if info["conductivity"] else "")
            + (f'<p><b>Fraction mass% alignment:</b> {metric["mass_percent_observations"]} observations. '
               f'<a href="#mass-preview-run-{html.escape(str(run_number), quote=True)}">View mass% graph</a>.</p>'
               if comp_refs else
               '<p><b>Fraction mass% alignment:</b> no observations enabled for this run. '
               'Select mass% refinement mode, turn on this run\'s fraction mass%, and enter fraction values to include them.</p>')
            + _chromatogram_fit_svg(config, info)
            + (_conductivity_fit_svg(info) if info["conductivity"] else "")
            + (f'<h3 id="mass-preview-run-{html.escape(str(run_number), quote=True)}">Fraction mass% at midpoints</h3>'
               + _composition_profile_svg(config, comp_refs, info["result"]["trace"])
               + f"<p>{sum(len(row['mass_percent']) for row in comp_refs)} entered species observations. "
                 "Each populated cell is compared with integrated modeled mass over its collected interval; connecting curves are display-only.</p>"
               if comp_refs else "")
        ))
    default_run = str(next(
        (item["run_number"] for item in metrics if item["mass_percent_observations"]),
        metrics[0]["run_number"],
    ))
    choices = [
        f'<option value="{html.escape(str(item["run_number"]), quote=True)}"'
        + (' selected' if str(item["run_number"]) == default_run else '')
        + f'>LSQ Run {html.escape(str(item["run_number"]))}: {html.escape(item["run_name"])}'
        + (f' — {item["mass_percent_observations"]} mass% points' if item["mass_percent_observations"] else '')
        + '</option>'
        for item in metrics
    ]
    rendered_sections = ''.join(
        f'<section class="lsq-preview-run" data-run="{html.escape(run_number, quote=True)}"'
        + (' hidden>' if run_number != default_run else '>') + content + '</section>'
        for run_number, content in sections
    )
    output_path = Path(output_html)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        '<!doctype html><meta charset="utf-8"><title>Pre-refinement alignment</title>'
        '<style>body{font-family:Arial,sans-serif;margin:2rem;max-width:1100px}'
        'svg{display:block;width:100%;height:auto}h2{margin-top:2rem}'
        '[hidden]{display:none!important}select{font-size:1rem;padding:.4rem;margin:.4rem}</style>'
        '<h1>Measured versus model alignment before refinement</h1>'
        '<p>These are the current model parameters and selected saved model-run recipes. '
        'No optimizer was run and no fitted parameters were changed. Measured UV is displayed '
        'as the median imported sample in each chart pixel column to prevent dark overdraw; '
        'every imported sample still enters the least-squares calculation. Predicted UV spans the complete programmed run. '
        'The selected detector-baseline policy is applied exactly as it is during refinement. '
        'If fraction mass% is entered, the preview plots observations at fraction midpoints and compares them with modeled mass integrated over each full collected interval. Blank cells add no residual.</p>'
        '<label for="lsq-preview-run-picker">View LSQ run:</label>'
        '<select id="lsq-preview-run-picker">' + ''.join(choices) + '</select>'
        + rendered_sections
        + '<script>const picker=document.getElementById("lsq-preview-run-picker");'
          'function showRun(){document.querySelectorAll(".lsq-preview-run").forEach(section=>'
          '{section.hidden=section.dataset.run!==picker.value;});}'
          'picker.addEventListener("change",showRun);showRun();</script>', encoding="utf-8",
    )
    return {"report_html": str(output_path), "profiles": metrics}


def _predicted_mass_percent_at_cv(sim_cv: np.ndarray, comp: np.ndarray,
                                  cv: float, species_count: int) -> np.ndarray:
    concentrations = np.array([np.interp(cv, sim_cv, comp[:, i])
                               for i in range(species_count)], dtype=float)
    total = float(np.sum(concentrations))
    return 100.0 * (concentrations / total) if total > 1e-12 else np.zeros(species_count)


def _predicted_fraction_mass_percent(config: dict[str, Any], trace: dict[str, Any],
                                     ref: dict[str, Any]) -> tuple[np.ndarray, bool]:
    """Integrate detector species concentration over the collected volume.

    CV is cumulative volume / column volume, so Q dt is exactly V_col dCV,
    including when the stage flow changes. The constant V_col cancels in the
    mass percentages. Interpolated endpoints include partial simulation cells.
    Legacy point observations retain their explicitly pointwise meaning.
    """
    cv = np.asarray(trace["CV"], dtype=float)
    comp = np.maximum(np.asarray(trace["component_g_L"], dtype=float), 0.0)
    if ref.get("legacy_point"):
        return _predicted_mass_percent_at_cv(cv, comp, float(ref["CV"]), comp.shape[1]), True
    factor = float(config["column"]["volume_mL"]) if ref["x_unit"] == "ML" else 1.0
    start, end = float(ref["start"]) / factor, float(ref["end"]) / factor
    tolerance = 1e-9 * max(1.0, float(cv[-1]))
    if start < cv[0] - tolerance or end > cv[-1] + tolerance:
        raise ValueError(f"Fraction {ref['fraction_index']} [{start:g}, {end:g}] CV lies outside the simulated run.")
    start, end = max(start, float(cv[0])), min(end, float(cv[-1]))
    points = np.r_[start, cv[(cv > start) & (cv < end)], end]
    concentrations = np.column_stack([np.interp(points, cv, comp[:, i])
                                      for i in range(comp.shape[1])])
    integrated = np.sum(0.5 * (concentrations[:-1] + concentrations[1:])
                        * np.diff(points)[:, None], axis=0)
    total = float(np.sum(integrated))
    quantifiable = total / max(end - start, np.finfo(float).tiny) > 1e-12
    return (100.0 * (integrated / total) if quantifiable else np.zeros(comp.shape[1])), quantifiable


def _composition_residuals(
    config: dict[str, Any],
    references: list[dict[str, Any]],
    trace: dict[str, Any],
) -> tuple[np.ndarray, list[dict[str, Any]]]:
    names = [str(r["name"]) for r in active_components(config)]
    sim_cv = np.asarray(trace["CV"], dtype=float)
    rows: list[dict[str, Any]] = []
    residuals: list[float] = []
    for ref in references:
        if not ref["mass_percent"]:
            continue
        cv = float(ref["CV"])
        if cv < float(np.min(sim_cv)) - 1e-12 or cv > float(np.max(sim_cv)) + 1e-12:
            raise ValueError(f"Composition reference at CV {cv:g} lies outside the simulated CV range.")
        predicted, quantifiable = _predicted_fraction_mass_percent(config, trace, ref)
        for i, name in enumerate(names):
            if name not in ref["mass_percent"]:
                continue
            measured = float(ref["mass_percent"][name])
            pred = float(predicted[i])
            residuals.append((measured - pred) / 100.0)
            rows.append({"CV": cv, "Fraction": ref["fraction_index"],
                         "Fraction_Start": ref["start"], "Fraction_End": ref["end"],
                         "Fraction_Unit": ref["x_unit"], "Protein_Species": name,
                         "Measured_mass_percent": measured, "Simulated_mass_percent": pred,
                         "Model_Fraction_Quantifiable": quantifiable})
    if residuals:
        arr = np.asarray(residuals, dtype=float) / math.sqrt(len(residuals))
    else:
        arr = np.empty(0, dtype=float)
    return arr, rows


def _composition_profile_svg(config: dict[str, Any], references: list[dict[str, Any]],
                             trace: dict[str, Any]) -> str:
    """Plot midpoint observations and model fractions without bridging data gaps.

    The smooth observed curve is display-only. The least-squares residual uses
    each entered midpoint exactly once; missing species cells add no residual.
    """
    names = [str(row["name"]) for row in active_components(config)]
    colors = ("#0878c6", "#e87722", "#149447", "#dc3545", "#8a63b7", "#9b6a42")
    sim_cv = np.asarray(trace["CV"], dtype=float)
    show_ml = all(ref["x_unit"] == "ML" for ref in references)
    x_factor = float(config["column"]["volume_mL"]) if show_ml else 1.0
    x_min = 0.0
    x_max = max(float(sim_cv[-1]), max(float(ref["CV"]) for ref in references), 1e-6) * x_factor
    left, top, width, height = 75.0, 30.0, 680.0, 345.0
    px = lambda cv: left + width * (float(cv) * x_factor - x_min) / (x_max - x_min)
    py = lambda pct: top + height * (2.0 - math.log10(max(float(pct), 0.01))) / 4.0

    def contiguous(a: dict[str, Any], b: dict[str, Any]) -> bool:
        if int(b["fraction_index"]) != int(a["fraction_index"]) + 1:
            return False
        if a.get("legacy_point") or b.get("legacy_point"):
            return True
        if a["x_unit"] != b["x_unit"]:
            return False
        return float(b["start"]) <= float(a["end"]) + 1e-6 * max(1.0, float(a["end"]))

    def path(points: list[tuple[float, float]]) -> str:
        return " ".join(("M" if i == 0 else "L") + f"{px(cv):.2f},{py(value):.2f}"
                        for i, (cv, value) in enumerate(points))

    svg = ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 960 440" '
           'role="img" aria-label="Fraction midpoint mass percent compared with model; missing values leave gaps">',
           '<rect width="960" height="440" fill="white"/>']
    for tick in (0.01, 0.1, 1.0, 10.0, 100.0):
        y = py(tick)
        svg.append(f'<line x1="{left}" y1="{y:.2f}" x2="{left+width}" y2="{y:.2f}" stroke="#e2e6ea"/>')
        svg.append(f'<text x="{left-12}" y="{y+4:.2f}" text-anchor="end" font-size="12">{tick:g}</text>')
    for i in range(6):
        display_value = x_min + (x_max - x_min) * i / 5
        x = left + width * i / 5
        svg.append(f'<line x1="{x:.2f}" y1="{top}" x2="{x:.2f}" y2="{top+height}" stroke="#edf0f2"/>')
        svg.append(f'<text x="{x:.2f}" y="{top+height+20}" text-anchor="middle" font-size="12">{display_value:.3g}</text>')
    x_label = "mL" if show_ml else "CV"
    svg.append(f'<text x="{left+width/2}" y="425" text-anchor="middle" font-size="14">Fraction midpoint [{x_label} from run start]</text>')
    svg.append('<text x="18" y="203" transform="rotate(-90 18 203)" text-anchor="middle" font-size="14">Mass% (log scale; zero at floor)</text>')

    for species_index, name in enumerate(names):
        color = colors[species_index]
        observed = [ref for ref in references if name in ref["mass_percent"]]
        groups: list[list[dict[str, Any]]] = []
        for ref in observed:
            if not groups or not contiguous(groups[-1][-1], ref):
                groups.append([])
            groups[-1].append(ref)
        for group in groups:
            if len(group) > 1:
                xs = np.asarray([float(ref["CV"]) for ref in group])
                ys = np.asarray([float(ref["mass_percent"][name]) for ref in group])
                smooth_x = np.linspace(xs[0], xs[-1], max(24, 14 * (len(group) - 1)))
                smooth_y = np.clip(PchipInterpolator(xs, ys)(smooth_x), 0.0, 100.0)
                measured_style = ' stroke-dasharray="2 4"' if species_index else ''
                svg.append(f'<path d="{path(list(zip(smooth_x, smooth_y)))}" fill="none" stroke="{color}" stroke-width="2.3"{measured_style}/>')
                # Display fraction-integrated predictions at the actual
                # collection intervals. Interpolation is visual only.
                model_x = xs
                model_y = [_predicted_fraction_mass_percent(config, trace, ref)[0][species_index]
                           for ref in group]
                model_pattern = "1 6" if species_index else "6 4"
                model_opacity = ".6" if species_index else ".85"
                svg.append(f'<path d="{path(list(zip(model_x, model_y)))}" fill="none" stroke="{color}" stroke-width="1.8" stroke-dasharray="{model_pattern}" opacity="{model_opacity}"/>')
            for ref in group:
                cv = float(ref["CV"])
                measured = float(ref["mass_percent"][name])
                predicted = float(_predicted_fraction_mass_percent(config, trace, ref)[0][species_index])
                svg.append(f'<circle cx="{px(cv):.2f}" cy="{py(measured):.2f}" r="3.4" fill="{color}"/>')
                svg.append(f'<circle cx="{px(cv):.2f}" cy="{py(predicted):.2f}" r="3" fill="white" stroke="{color}" stroke-width="1.5"/>')
        legend_y = 55 + species_index * 25
        legend_style = ' stroke-dasharray="2 4"' if species_index else ''
        svg.append(f'<line x1="790" y1="{legend_y}" x2="813" y2="{legend_y}" stroke="{color}" stroke-width="3"{legend_style}/>')
        svg.append(f'<text x="820" y="{legend_y+4}" font-size="12">{html.escape(name)}</text>')
    svg.append('<text x="790" y="245" font-size="12">Solid: measured target</text>')
    svg.append('<text x="790" y="264" font-size="12">Dotted: measured impurities</text>')
    svg.append('<text x="790" y="283" font-size="12">Dashed: model target</text>')
    svg.append('<text x="790" y="302" font-size="12">Faint dots: model impurities</text>')
    svg.append('<text x="790" y="321" font-size="12">Blank: no fit or line</text>')
    svg.append('</svg>')
    return "".join(svg)


def _fraction_uv_overlay_svg(config: dict[str, Any], references: list[dict[str, Any]],
                             info: dict[str, Any]) -> str:
    """One chromatogram with collected fraction percentages on a second axis."""
    trace = info["result"]["trace"]
    chrom = info["chromatogram"]
    program = info["result"]["simulation"]["program"]
    x_max = max(float(program.total_CV), 1e-9)
    measured_x, measured_y = _reference_display_series(chrom["CV"], chrom["measured"], x_max)
    model_x = np.asarray(trace["CV"], dtype=float)
    model_y = np.asarray(trace["uv_mAU"], dtype=float) + float(chrom["baseline"])
    y_min = min(0.0, float(np.min(measured_y)), float(np.min(model_y)))
    y_max = max(1.0, float(np.max(measured_y)), float(np.max(model_y)))
    if y_min == y_max:
        y_max += 1.0
    left, right, top, bottom = 70.0, 760.0, 35.0, 365.0
    px = lambda value: left + (right - left) * np.clip(value / x_max, 0.0, 1.0)
    py_uv = lambda value: bottom - (bottom - top) * (value - y_min) / (y_max - y_min)
    py_pct = lambda value: bottom - (bottom - top) * np.clip(value / 100.0, 0.0, 1.0)

    def line(x: np.ndarray, y: np.ndarray, y_map: Any) -> str:
        return " ".join(("M" if i == 0 else "L") + f"{px(float(a)):.2f},{y_map(float(b)):.2f}"
                        for i, (a, b) in enumerate(zip(x, y)))

    parts = ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1030 430" '
             'role="img" aria-label="Measured and predicted UV chromatogram with fraction mass percentages">',
             '<rect width="1030" height="430" fill="white"/>']
    for stage in _stage_receipt(config, program):
        x = px(float(stage["start_CV"]))
        parts.append(f'<line x1="{x:.2f}" y1="{top}" x2="{x:.2f}" y2="{bottom}" stroke="#d5d9de" stroke-dasharray="3 4"/>')
    for j in range(5):
        y = top + (bottom - top) * j / 4
        uv = y_max - (y_max - y_min) * j / 4
        pct = 100 - 25 * j
        parts.extend([f'<line x1="{left}" y1="{y:.2f}" x2="{right}" y2="{y:.2f}" stroke="#e4e7eb"/>',
                      f'<text x="{left-8}" y="{y+4:.2f}" text-anchor="end" font-size="11">{uv:.3g}</text>',
                      f'<text x="{right+8}" y="{y+4:.2f}" font-size="11">{pct}</text>'])
    for j in range(6):
        x = left + (right - left) * j / 5
        parts.append(f'<text x="{x:.2f}" y="{bottom+19}" text-anchor="middle" font-size="11">{x_max*j/5:.3g}</text>')
    parts.append(f'<path d="{line(measured_x, measured_y, py_uv)}" fill="none" stroke="#20262e" stroke-width="1.7"/>')
    parts.append(f'<path d="{line(model_x, model_y, py_uv)}" fill="none" stroke="#0878c6" stroke-width="1.8"/>')
    colors = ("#a32673", "#e87722", "#149447", "#dc3545", "#8a63b7", "#9b6a42")
    names = [str(row["name"]) for row in active_components(config)]
    for i, name in enumerate(names):
        groups: list[list[dict[str, Any]]] = []
        for ref in references:
            if name not in ref["mass_percent"]:
                continue
            if not groups or ref["fraction_index"] != groups[-1][-1]["fraction_index"] + 1 or (
                not ref.get("legacy_point") and not groups[-1][-1].get("legacy_point") and
                float(ref["start"]) > float(groups[-1][-1]["end"]) + 1e-8):
                groups.append([])
            groups[-1].append(ref)
        for group in groups:
            x = np.asarray([float(ref["CV"]) for ref in group])
            measured = np.asarray([float(ref["mass_percent"][name]) for ref in group])
            predicted = np.asarray([_predicted_fraction_mass_percent(config, trace, ref)[0][i]
                                    for ref in group])
            if len(group) > 1:
                parts.append(f'<path d="{line(x, measured, py_pct)}" fill="none" stroke="{colors[i]}" stroke-width="2"/>')
                parts.append(f'<path d="{line(x, predicted, py_pct)}" fill="none" stroke="{colors[i]}" stroke-width="1.7" stroke-dasharray="5 4"/>')
            for xv, mv, pv in zip(x, measured, predicted):
                parts.append(f'<circle cx="{px(xv):.2f}" cy="{py_pct(mv):.2f}" r="3" fill="{colors[i]}"/>')
                parts.append(f'<circle cx="{px(xv):.2f}" cy="{py_pct(pv):.2f}" r="3" fill="white" stroke="{colors[i]}"/>')
        y = 60 + 24 * i
        parts.append(f'<text x="805" y="{y}" fill="{colors[i]}" font-size="12">{html.escape(name)} mass%</text>')
    parts.extend([f'<text x="{(left+right)/2:.1f}" y="410" text-anchor="middle" font-size="13">CV from run start</text>',
                  '<text x="15" y="200" transform="rotate(-90 15 200)" font-size="12" text-anchor="middle">UV [mAU]</text>',
                  '<text x="800" y="210" font-size="12">Mass% (right axis)</text>',
                  '<text x="805" y="275" font-size="11">Black: measured UV</text>',
                  '<text x="805" y="293" font-size="11">Blue: predicted UV</text>',
                  '<text x="805" y="311" font-size="11">Filled: measured mass%</text>',
                  '<text x="805" y="329" font-size="11">Open/dashed: predicted mass%</text>',
                  '</svg>'])
    return "".join(parts)


def _objective_parts(
    cfg: dict[str, Any],
    ref: ChromatogramReference,
    mode: str,
    composition_refs: list[dict[str, Any]],
    fit_conductivity: bool = False,
    solver_deadline: float | None = None,
) -> tuple[np.ndarray, dict[str, Any]]:
    result = simulate(cfg, solver_deadline=solver_deadline) if solver_deadline is not None else simulate(cfg)
    chrom_r, chrom_info = _chromatogram_residuals(cfg, ref, result["trace"])
    cond_r = np.empty(0, dtype=float)
    cond_info = None
    if fit_conductivity:
        cond_r, cond_info = _conductivity_residuals(cfg, ref, result["trace"], chrom_info)
    comp_r = np.empty(0, dtype=float)
    comp_rows: list[dict[str, Any]] = []
    if mode == MODE_CHROMATOGRAM_AND_COMPOSITION:
        comp_r, comp_rows = _composition_residuals(cfg, composition_refs, result["trace"])
        if cfg["least_squares"]["objective"] == "RAW_SSE":
            # Keep the literal UV SSE, but express the dimensionless mass%
            # group in comparable UV units. Without this, mAU residuals can
            # drown out every fraction observation by orders of magnitude.
            comp_r = comp_r * max(1.0, float(chrom_info["measured_signal_scale"])) * math.sqrt(len(chrom_r))
        elif cfg["least_squares"]["objective"] == "NORMALIZED_MSE":
            # A flat/near-zero UV file must not divide by numerical noise and
            # overwhelm genuine fraction observations in the combined fit.
            chrom_r = chrom_r * min(1.0, float(chrom_info["measured_signal_scale"]))
    combined = np.concatenate([chrom_r, cond_r, comp_r])
    return combined, {
        "result": result,
        "chromatogram": chrom_info,
        "composition": comp_rows,
        "conductivity": cond_info,
        "conductivity_sse_normalized": float(np.sum(cond_r ** 2)),
        "chromatogram_sse_normalized": float(np.sum(chrom_r**2)),
        "chromatogram_raw_sse": float(np.sum((chrom_info["predicted_detector_signal"] - chrom_info["measured"])**2)),
        "composition_sse_normalized": float(np.sum(comp_r**2)) if len(comp_r) else 0.0,
    }


def _trial_deadline(config: dict[str, Any]) -> float | None:
    """Limit one TDM+CPA LSQ solve; normal predictions remain unrestricted."""
    minutes = float(config["least_squares"].get("max_trial_minutes", 0.0))
    return (time.perf_counter() + 60.0 * minutes
            if config["model"] == MODEL_TDM_CPA and minutes > 0 else None)


def _cpa_binding_metrics(initial_info: dict[str, Any], final_info: dict[str, Any]) -> dict[str, Any] | None:
    """Keep chromatographic validity separate from numerical fit acceptance."""
    before = initial_info["result"]["simulation"]["diagnostics"].get("CPA_adsorption_audit")
    after = final_info["result"]["simulation"]["diagnostics"].get("CPA_adsorption_audit")
    if not before or not after:
        return None
    peak = final_info["result"]["parameter_receipt"]["target_peak"]
    early_peak = (peak["column_outlet_peak_CV"] is not None
                  and peak["column_outlet_peak_CV"] < peak["first_elution_start_CV"] - 1e-9)
    return {
        "initial": before,
        "fitted": after,
        "target_peak_before_elution": bool(early_peak),
        "target_peak_stage": peak["stage_at_detector_peak"],
        "target_column_outlet_peak_CV": peak["column_outlet_peak_CV"],
        "target_peak_stage_from_adaptive_solver": after["species"][0]["peak_outlet_stage"],
        "first_elution_start_CV": peak["first_elution_start_CV"],
        "warnings": list(final_info["result"]["simulation"]["diagnostics"].get("warnings", [])),
    }


def _fit_acceptance(config: dict[str, Any], initial_obj: float, final_obj: float,
                    final_infos: Iterable[dict[str, Any]]) -> dict[str, Any]:
    """Require a measurable objective gain and, by default, bind-and-elute for TDM–CPA."""
    policy = config["least_squares"]
    gain_percent = _objective_improvement_percent(initial_obj, final_obj)
    minimum_gain = float(policy["minimum_objective_improvement_percent"])
    numerical = bool(math.isfinite(final_obj) and final_obj < initial_obj
                     and gain_percent >= minimum_gain)
    reasons: list[str] = []
    physical = True
    if config["model"] == MODEL_TDM_CPA and policy["physical_acceptance_mode"] == "REQUIRE_BIND_ELUTE":
        for run, info in enumerate(final_infos, start=1):
            audit = info["result"]["simulation"]["diagnostics"].get("CPA_adsorption_audit")
            if not audit or not audit["species"]:
                reasons.append(f"Run {run}: CPA binding audit is unavailable.")
                continue
            target = audit["species"][0]
            if target["end_load_bound_percent_of_load"] < float(policy["minimum_bound_percent_after_load"]):
                reasons.append(f"Run {run}: target bound after loading is below the selected minimum.")
            if target["peak_outlet_stage_kind"] != "ELUTION":
                reasons.append(f"Run {run}: target column-outlet maximum is in {target['peak_outlet_stage_kind']}, not elution.")
            if target["elution_recovery_percent_of_end_load_retained"] < float(policy["minimum_elution_recovery_percent"]):
                reasons.append(f"Run {run}: target recovery relative to end-load column inventory is below the selected minimum.")
        physical = not reasons
    status = ("ACCEPTED" if numerical and physical else
              "NUMERICALLY_IMPROVED_BUT_PHYSICALLY_INVALID" if numerical else "REJECTED")
    if not numerical:
        reasons.insert(0, f"Objective improvement {gain_percent:.6g}% does not exceed the required {minimum_gain:.6g}%.")
    return {"status": status, "accepted": status == "ACCEPTED",
            "numerically_accepted": numerical, "physically_accepted": physical,
            "reasons": reasons}


def _jacobian_diagnostics(opt: Any, selected_vector: np.ndarray,
                          specs: list[FitParameter]) -> dict[str, Any]:
    """Report sensitivity only when SciPy's Jacobian belongs to the selected fit."""
    result: dict[str, Any] = {
        "optimizer_jacobian_evaluations": (int(opt.njev) if getattr(opt, "njev", None) is not None else None),
        "optimizer_optimality": (float(opt.optimality) if getattr(opt, "optimality", None) is not None else None),
        "finite_difference_step_rule": "in transformed coordinates: max(0.001, abs(value)*0.001), one-sided at bounds",
        "rank": None, "singular_values": None, "condition_estimate": None,
        "sensitivity_by_parameter": None,
        "note": "Jacobian diagnostics unavailable at the selected best candidate.",
    }
    jac = getattr(opt, "jac", None)
    if jac is None or not np.allclose(np.asarray(getattr(opt, "x", []), dtype=float),
                                      selected_vector, rtol=0, atol=1e-12):
        return result
    matrix = np.asarray(jac.toarray() if hasattr(jac, "toarray") else jac, dtype=float)
    if matrix.ndim != 2 or matrix.shape[1] != len(specs) or not np.all(np.isfinite(matrix)):
        return result
    singular = np.linalg.svd(matrix, compute_uv=False)
    rank = int(np.linalg.matrix_rank(matrix))
    result.update(rank=rank, singular_values=singular.tolist(),
                  condition_estimate=(float(singular[0] / singular[-1]) if len(singular)
                                      and singular[-1] > 0 else None),
                  sensitivity_by_parameter={spec.path: float(np.linalg.norm(matrix[:, i]))
                                            for i, spec in enumerate(specs)},
                  note=("Jacobian is rank deficient; parameter uncertainty is not identifiable."
                        if rank < len(specs) else
                        "Local sensitivities use transformed fit coordinates. Confidence intervals require a justified noise model and are not reported."))
    return result


def _cpa_binding_html(metrics: dict[str, Any] | None) -> str:
    if not metrics:
        return ""
    initial = metrics["initial"]["species"]
    fitted = metrics["fitted"]["species"]
    rows = "".join(
        "<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in (
            after["name"], f'{before["zero_coverage_partition_load"]:.5g}',
            f'{after["zero_coverage_partition_load"]:.5g}',
            f'{before["end_load_bound_percent_of_load"]:.3g}',
            f'{after["end_load_bound_percent_of_load"]:.3g}',
            f'{after["load_outlet_percent_of_load"]:.3g}',
            f'{after["wash_outlet_percent_of_load"]:.3g}',
            f'{after["elution_outlet_percent_of_load"]:.3g}',
        )) + "</tr>" for before, after in zip(initial, fitted)
    )
    warning = " ".join(metrics["warnings"])
    warning_html = f'<p><b>Binding warning:</b> {html.escape(warning)}</p>' if warning else ""
    intro = (
        "<h4>CPA binding and elution audit</h4>"
        "<p>Zero-coverage K and adsorbed mass after loading show whether the fitted model predicts binding. "
        "Stage outlet percentages are integrated at the column before detector delay, divided by species load. "
        "A lower UV objective alone does not establish bind-and-elute behaviour.</p>"
        f'<p><b>Target column-outlet peak:</b> {html.escape(str(metrics["target_column_outlet_peak_CV"]))} CV; '
        f'first elution begins at {metrics["first_elution_start_CV"]:.5g} CV. '
        f'Adaptive-solver peak stage: {html.escape(metrics["target_peak_stage_from_adaptive_solver"])}; '
        f'elution recovery of end-load retained target: {fitted[0]["elution_recovery_percent_of_end_load_retained"]:.3g}%.</p>'
    )
    table = (
        "<table><thead><tr><th>Species</th><th>Initial K load</th><th>Fitted K load</th>"
        "<th>Initial bound after load %</th><th>Fitted bound after load %</th>"
        "<th>Fitted load outlet %</th><th>Fitted wash outlet %</th>"
        "<th>Fitted elution outlet %</th></tr></thead><tbody>" + rows + "</tbody></table>"
    )
    stage_rows = "".join(
        "<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in (
            stage["stage"], f'{stage["outlet_g"][0]:.5g}',
            f'{stage["bulk_inventory_g"][0]:.5g}',
            f'{stage["pore_inventory_g"][0]:.5g}',
            f'{stage["bound_inventory_g"][0]:.5g}',
            f'{stage["column_inventory_g"][0]:.5g}',
            f'{stage["cumulative_mass_balance_error_percent_of_load"][0]:.4g}',
        )) + "</tr>" for stage in metrics["fitted"]["stages"]
    )
    stage_table = ("<h5>Target process-stage audit</h5><table><thead><tr><th>Stage</th>"
                   "<th>Outlet [g]</th><th>Bulk [g]</th><th>Pore [g]</th>"
                   "<th>Bound [g]</th><th>Column inventory [g]</th>"
                   "<th>Cumulative mass closure [% of load]</th></tr></thead><tbody>"
                   + stage_rows + "</tbody></table>")
    return intro + warning_html + table + stage_table


class EvaluationBudgetReached(RuntimeError):
    """Stop optimization before the selected model-evaluation budget is exceeded."""


def _write_apply_jsl(config: dict[str, Any], specs: list[FitParameter]) -> None:
    mapping = {
        "qmax_g_L": "Qmax", "b_L_g": "B", "Z_ref": "Zref", "delta_ref": "Delta",
        "As_m_inv": "As", "kkin_star_s": "Kkin", "salt_sensitivity_per_M": "SaltSensitivity",
        "uv_response_factor_mAU_L_g": "UVResponse", "D_ax_mm2_s": "Dax",
        "k_eff_um_s": "Keff", "accessible_particle_porosity": "EpsPi",
        "D_app_mm2_s": "Dapp", "diameter_nm": "Diameter", "pH_ref": "PHref",
        "Z1_per_pH": "Z1", "Z2_per_pH2": "Z2", "Z3_per_pH3": "Z3",
        "delta_pH_slope_m2_C": "DeltaSlope",
        "wang_K_kin_s": "WangKkin", "wang_k_eq": "WangKeq",
        "wang_qmax_g_L": "WangQmax", "wang_n": "WangN",
        "wang_beta0": "WangBeta0", "wang_beta1_per_M": "WangBeta1",
        "wang_beta2_L_g": "WangBeta2", "wang_beta3_per_pH": "WangBeta3",
    }
    global_mapping = {
        "conversion.uv_to_protein_mAU_L_g": "uvToProtein",
        "conversion.conductivity_to_salt_M_per_mS_cm": "condToSalt",
        "tdm.bead_radius_um": "beadRadius",
        "tdm.void_fraction": "voidFraction",
        "tdm.particle_porosity": "particlePorosity",
        "tdm.salt_axial_dispersion_mm2_s": "saltDax",
        "edm.total_bed_porosity": "edmPorosity",
        "cpa.ligand_surface_density_umol_m2": "ligandSurface",
        "cpa.system_specific_adsorption_parameter": "systemAdsorption",
        "wang.q0_g_L": "wangQ0", "wang.eta": "wangEta",
        "system.buffer_dispersion_mL": "bufferDispersion",
        "system.salt_dispersion_mm2_s": "systemSaltDax",
        "system.uv_dead_volume_mL": "uvDeadVolume",
        "system.conductivity_dead_volume_mL": "condDeadVolume",
    }
    lines = ["// Generated by least-squares refinement; updates only fitted mechanistic parameter boxes."]
    for spec in specs:
        tokens = _parse_path(spec.path)
        if tokens[0] == "components":
            comp_index = int(tokens[1]) + 1
            field = str(tokens[-1])
            suffix = mapping[field]
            variable = f"c{comp_index}{suffix}"
        else:
            variable = global_mapping[spec.path]
        value = float(_get_path(config, spec.path))
        lines.append(f"{variable} << Set({value:.17g});")
    lines.append("UpdateUI();")
    FIT_APPLY_JSL.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _profile_output_path(path: Path, run_context: dict[str, Any] | None) -> Path:
    if not run_context or run_context.get("run_number") is None:
        return path
    try:
        run_number = int(run_context["run_number"])
    except (TypeError, ValueError):
        return path
    return path.with_name(f"{path.stem}_run_{run_number:02d}{path.suffix}")


def _objective_history_svg(history: list[dict[str, Any]]) -> str:
    """Create a dependency-free plot of the best objective versus evaluation."""
    width, height = 900, 300
    left, right, top, bottom = 82, 26, 24, 54
    plot_w, plot_h = width - left - right, height - top - bottom
    valid = [row for row in history if math.isfinite(float(row.get("best_objective", math.nan)))]
    if not valid:
        return ""
    xs = [float(row["evaluation"]) for row in valid]
    ys = [max(float(row["best_objective"]), 1e-16) for row in valid]
    xmin, xmax = min(xs), max(xs)
    if xmax <= xmin:
        xmax = xmin + 1.0
    log_values = [math.log10(v) for v in ys]
    ymin, ymax = min(log_values), max(log_values)
    if ymax - ymin < 0.15:
        ymin -= 0.075
        ymax += 0.075
    else:
        pad = 0.06 * (ymax - ymin)
        ymin -= pad
        ymax += pad

    def px(value: float) -> float:
        return left + (value - xmin) / (xmax - xmin) * plot_w

    def py(value: float) -> float:
        return top + (ymax - math.log10(max(value, 1e-16))) / (ymax - ymin) * plot_h

    points = " ".join(f"{px(x):.2f},{py(y):.2f}" for x, y in zip(xs, ys))
    grid = []
    for j in range(4):
        y = top + j * plot_h / 3
        value = 10 ** (ymax - j * (ymax - ymin) / 3)
        grid.append(
            f'<line x1="{left}" y1="{y:.2f}" x2="{left+plot_w}" y2="{y:.2f}" stroke="#e1e5e9"/>'
            f'<text x="{left-8}" y="{y+4:.2f}" text-anchor="end" font-size="12">{value:.3g}</text>'
        )
    return (
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" width="100%" role="img" '
        'aria-label="Best least-squares objective by evaluation, lower is better">'
        f'<rect width="{width}" height="{height}" fill="white"/>{"".join(grid)}'
        f'<polyline points="{points}" fill="none" stroke="#1769aa" stroke-width="2.5"/>'
        f'<text x="{left+plot_w/2:.1f}" y="{height-12}" text-anchor="middle" font-size="13">Residual evaluation</text>'
        f'<text transform="translate(18 {top+plot_h/2:.1f}) rotate(-90)" text-anchor="middle" font-size="13">Best objective (log scale)</text>'
        f'<text x="{left}" y="{height-30}" text-anchor="start" font-size="11">{int(xmin)}</text>'
        f'<text x="{left+plot_w}" y="{height-30}" text-anchor="end" font-size="11">{int(max(xs))}</text>'
        '</svg>'
    )


def _objective_improvement_percent(initial: float, candidate: float) -> float:
    """Return a stable percentage change, treating machine-scale objectives as zero."""
    initial = float(initial)
    candidate = float(candidate)
    if abs(initial) <= 1e-12 and abs(candidate - initial) <= 1e-12:
        return 0.0
    return (initial - candidate) / max(abs(initial), 1e-30) * 100.0


def _format_elapsed_hms(seconds: float) -> str:
    whole_seconds = max(0, round(seconds))
    hours, remainder = divmod(whole_seconds, 3600)
    minutes, seconds_left = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds_left:02d}"


def _bounded_jacobian(fun, lower, upper, batch_fun=None, trial_valid=None,
                      trial_timed_out=None, parameter_labels=None, probe_callback=None):
    def jacobian(vector):
        base = fun(vector)
        if trial_valid is not None and not trial_valid(vector):
            if trial_timed_out is not None and trial_timed_out(vector):
                raise TimeoutError("TDM+CPA LSQ timed out at the current parameter set; increase the trial timeout or revise the starting parameters.")
            return np.zeros((len(base), len(vector)), dtype=float)
        trials = []
        steps = []
        for i, value in enumerate(vector):
            h = max(1e-3, abs(float(value)) * 1e-3)
            trial = np.asarray(vector, dtype=float).copy()
            step = min(h, upper[i] - value)
            if step < h * 0.1:
                step = -min(h, value - lower[i])
            trial[i] += step
            trials.append(trial)
            steps.append(step)
        values = batch_fun(trials) if batch_fun is not None else [fun(trial) for trial in trials]
        columns = []
        for i, (trial, result, step) in enumerate(zip(trials, values, steps)):
            if trial_valid is not None and not trial_valid(trial):
                if trial_timed_out is not None and trial_timed_out(trial):
                    # A timed-out direction is not a numerical derivative.
                    # Try the opposite direction once; if that also fails,
                    # stop with the exact parameter named instead of silently
                    # setting a zero Jacobian column or waiting for hours.
                    reverse_step = -step
                    if lower[i] <= vector[i] + reverse_step <= upper[i]:
                        reverse = np.asarray(vector, dtype=float).copy()
                        reverse[i] += reverse_step
                        if probe_callback is not None:
                            probe_callback(i)
                        reverse_result = (batch_fun([reverse])[0]
                                          if batch_fun is not None else fun(reverse))
                        if trial_valid(reverse):
                            columns.append((reverse_result - base) / reverse_step)
                            continue
                    label = (parameter_labels[i] if parameter_labels is not None
                             else f"parameter {i + 1}")
                    raise TimeoutError(
                        f"TDM+CPA LSQ timed out probing {label} in the feasible directions. "
                        "Lock this parameter or increase the trial timeout, then rerun the fit.")
                # An invalid CPA trajectory is a feasibility boundary, not a
                # measured residual. Never use its fixed penalty to form a
                # numerical derivative; probe the other feasible direction
                # and then reduce the step if necessary.
                h = abs(step)
                for scale in (1.0, 0.5, 0.25, 0.125, 0.0625):
                    for direction in (-np.sign(step), np.sign(step)):
                        if scale == 1.0 and direction == np.sign(step):
                            continue  # This is the trial already known to be invalid.
                        alternative_step = float(direction * h * scale)
                        if not lower[i] <= vector[i] + alternative_step <= upper[i]:
                            continue
                        alternative = np.asarray(vector, dtype=float).copy()
                        alternative[i] += alternative_step
                        if probe_callback is not None:
                            probe_callback(i)
                        alternative_result = (batch_fun([alternative])[0]
                                              if batch_fun is not None else fun(alternative))
                        if trial_valid(alternative):
                            result, step = alternative_result, alternative_step
                            break
                    else:
                        continue
                    break
                else:
                    columns.append(np.zeros_like(base))
                    continue
            columns.append((result - base) / step)
        return np.column_stack(columns)
    return jacobian


def _cache_last_evaluation(fun):
    """Reuse the residual SciPy just evaluated at the Jacobian base point."""
    last_vector = None
    last_result = None

    def cached(vector):
        nonlocal last_vector, last_result
        vector = np.asarray(vector, dtype=float)
        if last_vector is not None and np.array_equal(vector, last_vector):
            return last_result.copy()
        result = np.asarray(fun(vector), dtype=float)
        last_vector = vector.copy()
        last_result = result.copy()
        return result

    return cached


def _optimizer_method(initial: np.ndarray, lower: np.ndarray, upper: np.ndarray) -> str:
    """A zero-valued dispersion parameter must be able to leave its bound.

    TRF starts infinitesimally inside an active bound and can make only tiny
    steps for many evaluations. Dogbox takes a finite feasible step from the
    entered value; use TRF for the usual interior starting points.
    """
    return "dogbox" if np.any((initial <= lower) | (initial >= upper)) else "trf"


def run_least_squares_refinement(
    config: dict[str, Any],
    *,
    mode: str,
    chromatogram_csv: Path | str = ATTACHED_CHROMATOGRAM_CSV,
    composition_references: Iterable[dict[str, Any]] | None = None,
    run_context: dict[str, Any] | None = None,
    max_nfev: int = 60,
    unlocked_paths: Iterable[str] | None = None,
) -> dict[str, Any]:
    # Canonicalize legacy configs before validation, fitting and writing the
    # parameter receipt. This fills newly added impurity UV factors from the
    # shared target response factor and prevents a None initial value in reports.
    config = normalize_config(config)
    mode = str(mode or "").strip().upper()
    if mode not in SUPPORTED_MODES:
        raise ValueError("Least-squares mode must be CHROMATOGRAM or CHROMATOGRAM_AND_COMPOSITION.")
    errors = validate_config(config, strict=True)
    if errors:
        raise ValueError("Model inputs must be complete before refinement:\n- " + "\n- ".join(errors))
    if mode == MODE_CHROMATOGRAM_AND_COMPOSITION and config["least_squares"]["objective"] == "WEIGHTED_RMSE":
        raise ValueError("WEIGHTED_RMSE is defined for chromatogram signals only. For species mass% constraints select RAW_SSE or NORMALIZED_MSE.")
    fit_results_path = _profile_output_path(FIT_RESULTS_CSV, run_context)
    fit_trace_path = _profile_output_path(FIT_TRACE_CSV, run_context)
    fit_composition_path = _profile_output_path(FIT_COMPOSITION_CSV, run_context)
    fit_report_path = _profile_output_path(FIT_REPORT_HTML, run_context)
    fit_config_path = _profile_output_path(FIT_CONFIG_JSON, run_context)
    fit_settings_path = _profile_output_path(FIT_SETTINGS_JSON, run_context)
    fit_history_path = _profile_output_path(FIT_HISTORY_CSV, run_context)
    fit_apply_profile_path = _profile_output_path(FIT_APPLY_JSL, run_context)
    ref = load_chromatogram_reference(chromatogram_csv, config, require_uv=True)
    comp_refs: list[dict[str, Any]] = []
    if mode == MODE_CHROMATOGRAM_AND_COMPOSITION:
        comp_refs = validate_composition_references(config, composition_references or [])

    if unlocked_paths is not None:
        unlocked_paths = list(unlocked_paths)
        config["least_squares"]["unlocked_paths"] = list(unlocked_paths)
    specs = fit_parameter_specs(config, unlocked_paths=unlocked_paths)
    x0, lo, hi = _encode_initial(specs)
    fit_started = time.perf_counter()
    initial_r, initial_info = _objective_parts(config, ref, mode, comp_refs,
                                              solver_deadline=_trial_deadline(config))
    initial_obj = float(np.sum(initial_r**2))
    best_objective = initial_obj
    best_vector = None  # None preserves the original config without a transform round-trip.
    history: list[dict[str, Any]] = [{
        "evaluation": 0, "elapsed_seconds": time.perf_counter() - fit_started,
        "objective": initial_obj, "best_objective": initial_obj,
        "improvement_percent": 0.0, "phase": "initial", "status": "ok",
    }]
    eval_count = 0
    invalid_trial_vectors: set[bytes] = set()
    timed_out_trial_vectors: set[bytes] = set()

    def record_objective(objective: float, phase: str, status: str) -> None:
        nonlocal best_objective
        best_objective = min(best_objective, objective)
        improvement = _objective_improvement_percent(initial_obj, best_objective)
        history.append({
            "evaluation": eval_count,
            "elapsed_seconds": time.perf_counter() - fit_started,
            "objective": objective,
            "best_objective": best_objective,
            "improvement_percent": improvement,
            "phase": phase,
            "status": status,
        })

    def residual(vector: np.ndarray) -> np.ndarray:
        nonlocal eval_count, best_vector
        eval_count += 1
        try:
            candidate = _apply_vector(config, specs, vector)
            r, info = _objective_parts(candidate, ref, mode, comp_refs,
                                       solver_deadline=_trial_deadline(candidate))
            objective = float(np.sum(r*r))
            if not np.all(np.isfinite(r)):
                raise FloatingPointError("non-finite residual")
            if not math.isfinite(objective):
                raise FloatingPointError("non-finite objective")
            invalid_trial_vectors.discard(np.asarray(vector, dtype=float).tobytes())
            timed_out_trial_vectors.discard(np.asarray(vector, dtype=float).tobytes())
            if objective < best_objective:
                best_vector = np.asarray(vector).copy()
            record_objective(objective, "optimizer", "ok")
            return r
        except Exception as exc:
            # Keep the optimizer within the physically/numerically feasible region.
            invalid_trial_vectors.add(np.asarray(vector, dtype=float).tobytes())
            if isinstance(exc, TimeoutError):
                timed_out_trial_vectors.add(np.asarray(vector, dtype=float).tobytes())
            penalty = np.full_like(initial_r, 1e6, dtype=float)
            record_objective(float(np.sum(penalty * penalty)), "optimizer", f"penalty: {type(exc).__name__}: {exc}")
            return penalty

    if specs:
        cached_residual = _cache_last_evaluation(residual)
        opt = least_squares(
            cached_residual,
            x0,
            bounds=(lo, hi),
            jac=_bounded_jacobian(
                cached_residual, lo, hi,
                trial_valid=lambda vector: np.asarray(vector, dtype=float).tobytes()
                not in invalid_trial_vectors,
                trial_timed_out=lambda vector: np.asarray(vector, dtype=float).tobytes()
                in timed_out_trial_vectors,
                parameter_labels=[spec.label for spec in specs]),
            method=_optimizer_method(x0, lo, hi),
            loss="linear",
            x_scale="jac",
            # The model is integrated numerically at finite tolerance.  SciPy's
            # default machine-epsilon step can be smaller than the ODE noise and
            # produce a zero/unstable Jacobian, causing premature `xtol` stops.
            ftol=1e-4,
            xtol=1e-4,
            gtol=1e-7,
            max_nfev=max(2, int(max_nfev)),
        )
        jac_diagnostics = _jacobian_diagnostics(opt, best_vector if best_vector is not None else x0, specs)
    else:
        opt = SimpleNamespace(x=x0, success=True, message="All eligible parameters are locked; evaluated without refinement.", nfev=0)
        jac_diagnostics = _jacobian_diagnostics(opt, x0, specs)

    if best_vector is None:
        fitted = copy.deepcopy(config)
        opt.x = np.asarray([math.log10(float(_get_path(config, spec.path))) if spec.transform == "LOG10"
                            else float(_get_path(config, spec.path)) for spec in specs])
    else:
        opt.x = best_vector.copy()
        fitted = _apply_vector(config, specs, opt.x)
    final_r, final_info = _objective_parts(fitted, ref, mode, comp_refs,
                                          solver_deadline=_trial_deadline(fitted))
    final_obj = float(np.sum(final_r**2))
    eval_count += 1
    record_objective(final_obj, "final verification", "ok")

    # SciPy can stop at the cap after finding a useful candidate. Apply it only
    # when the verified objective and enabled physical checks both pass.
    acceptance = _fit_acceptance(config, initial_obj, final_obj, [final_info])
    accepted = acceptance["accepted"]

    # Keep the candidate as an audit receipt even when the active model is unchanged.
    save_config(fitted, fit_config_path)
    if accepted:
        save_config(fitted, CONFIG_PATH)
        _write_apply_jsl(fitted, specs)
    elif FIT_APPLY_JSL.exists():
        FIT_APPLY_JSL.unlink()
    if not accepted and fit_apply_profile_path.exists():
        fit_apply_profile_path.unlink()

    with fit_results_path.open("w", encoding="utf-8", newline="") as fh:
        fields = ["Parameter_Path", "Parameter", "Initial_Value", "Fitted_Value", "Transform",
                  "Fit_Range_Percent", "Allowed_Min", "Allowed_Max", "At_Lower_Bound", "At_Upper_Bound"]
        writer = csv.DictWriter(fh, fieldnames=fields); writer.writeheader()
        for spec, xfit in zip(specs, opt.x):
            initial_value = float(_get_path(config, spec.path))
            fitted_value = float(_get_path(fitted, spec.path))
            allowed_min, allowed_max = _physical_fit_bounds(spec)
            writer.writerow({
                "Parameter_Path": spec.path,
                "Parameter": spec.label,
                "Initial_Value": initial_value,
                "Fitted_Value": fitted_value,
                "Transform": spec.transform,
                "Fit_Range_Percent": config["least_squares"].get("fit_range_percent", {}).get(spec.path, ""),
                "Allowed_Min": allowed_min, "Allowed_Max": allowed_max,
                "At_Lower_Bound": abs(float(xfit) - spec.lower) <= 1e-6 * max(1.0, abs(spec.lower)),
                "At_Upper_Bound": abs(float(xfit) - spec.upper) <= 1e-6 * max(1.0, abs(spec.upper)),
            })

    chrom = final_info["chromatogram"]
    with fit_trace_path.open("w", encoding="utf-8", newline="") as fh:
        fields = ["CV", "Measured_Signal", "Simulated_Total_Protein_g_L", "Scaled_Simulated_Signal", "Residual"]
        writer = csv.DictWriter(fh, fieldnames=fields); writer.writeheader()
        for cv, ym, ys, yp in zip(chrom["CV"], chrom["measured"], chrom["simulated_total_g_L"], chrom["predicted_detector_signal"]):
            writer.writerow({"CV": cv, "Measured_Signal": ym, "Simulated_Total_Protein_g_L": ys, "Scaled_Simulated_Signal": yp, "Residual": float(yp)-float(ym)})

    with fit_composition_path.open("w", encoding="utf-8", newline="") as fh:
        fields = ["CV", "Fraction", "Fraction_Start", "Fraction_End", "Fraction_Unit", "Protein_Species", "Measured_mass_percent", "Simulated_mass_percent", "Model_Fraction_Quantifiable", "Residual_mass_percent"]
        writer = csv.DictWriter(fh, fieldnames=fields); writer.writeheader()
        for row in final_info["composition"]:
            r = dict(row)
            r["Residual_mass_percent"] = float(r["Measured_mass_percent"]) - float(r["Simulated_mass_percent"])
            writer.writerow(r)

    with fit_history_path.open("w", encoding="utf-8", newline="") as fh:
        history_fields = ["evaluation", "elapsed_seconds", "objective", "best_objective", "improvement_percent", "phase", "status"]
        writer = csv.DictWriter(fh, fieldnames=history_fields)
        writer.writeheader()
        writer.writerows(history)

    binding_metrics = _cpa_binding_metrics(initial_info, final_info)
    settings = {
        "mode": mode,
        "objective": config["least_squares"]["objective"],
        "baseline_mode": config["least_squares"]["baseline_mode"],
        "fitted_parameter_paths": [s.path for s in specs],
        "equation_9": "sum_points (predicted - measured)^2; only unlocked eligible parameters vary",
        "reference_csv": str(Path(chromatogram_csv).resolve()),
        "reference_x_column": ref.x_column,
        "reference_signal_column": ref.signal_column,
        "reference_x_unit": ref.x_unit,
        "reference_x_axis_assumption": ref.x_axis_assumption,
        "composition_references": comp_refs,
        "fitted_parameter_paths": [s.path for s in specs],
        "fit_range_percent": config["least_squares"].get("fit_range_percent", {}),
        "max_nfev": int(max_nfev),
        "objective_history_csv": str(fit_history_path),
        "target_run_context": run_context,
        "CPA_binding_audit": binding_metrics,
        "acceptance": acceptance,
        "jacobian_diagnostics": jac_diagnostics,
    }
    fit_settings_path.write_text(json.dumps(settings, indent=2, ensure_ascii=False), encoding="utf-8")

    parameter_rows = "".join(
        f"<tr><td>{html.escape(spec.label)}</td><td>{float(_get_path(config, spec.path)):.7g}</td><td>{float(_get_path(fitted, spec.path)):.7g}</td>"
        f"<td>{html.escape(str(config['least_squares'].get('fit_range_percent', {}).get(spec.path, 'Default')))}</td>"
        f"<td>{_physical_fit_bounds(spec)[0]:.7g}</td><td>{_physical_fit_bounds(spec)[1]:.7g}</td></tr>"
        for spec in specs
    )
    comp_note = (
        f"<p><b>Composition constraints:</b> {len(comp_refs)} fractions, {len(final_info['composition'])} populated species observations. "
        f"Composition objective contribution = {final_info['composition_sse_normalized']:.6g}.</p>"
        if mode == MODE_CHROMATOGRAM_AND_COMPOSITION else
        "<p><b>Composition constraints:</b> not used in this refinement.</p>"
    )
    x_assumption_note = (
        f"<p><b>X-axis assumption:</b> {html.escape(ref.x_axis_assumption)}</p>"
        if ref.x_axis_assumption else ""
    )
    run_context_html = (
        "<h2>Fit target run and operating recipe</h2>"
        "<p>The selected run's column, feed, numerical settings, buffers, load chemistry, "
        "PLW steps, and elution steps were held fixed while adsorption parameters were fitted.</p>"
        f"<details open><summary>Run recipe snapshot</summary><pre>{html.escape(json.dumps(run_context, indent=2, ensure_ascii=False))}</pre></details>"
        if run_context is not None else
        "<p><b>Fit target run:</b> not supplied by direct Python call.</p>"
    )
    improvement_percent = _objective_improvement_percent(initial_obj, final_obj)
    history_plot = _objective_history_svg(history)
    chromatogram_plot = _chromatogram_fit_svg(fitted, final_info)
    fraction_uv_plot = ("<h2>UV chromatogram with collected fraction mass%</h2>"
                        + _fraction_uv_overlay_svg(fitted, comp_refs, final_info)
                        if comp_refs else "")
    composition_plot = ("<h2>Fraction mass% profiles</h2>"
                        + _composition_profile_svg(fitted, comp_refs, final_info["result"]["trace"])
                        + "<p>Solid target and dotted impurity curves connect adjacent entered midpoints without overshoot. Model target is dashed; model impurities use lighter dots. Model lines appear only where corresponding measured segments exist. Blank cells create gaps and add no residual.</p>"
                        if comp_refs else "")
    uv_identifiability_note = (
        "<p>Impurity UV factors are estimated from the composite detector trace. When peaks overlap, their factors can be correlated with each other and with adsorption parameters; species-resolved UV data or composition constraints improve identifiability.</p>"
        if any(spec.path.endswith("uv_response_factor_mAU_L_g") for spec in specs) else
        "<p>Impurity UV factors were held fixed by the lock selection or were ineligible for the selected model. Only eligible unlocked response factors are refined.</p>"
    )
    fit_report_path.write_text(f"""<!doctype html>
<meta charset="utf-8"><title>Least-squares refinement</title>
<style>body{{font-family:Arial,sans-serif;margin:2rem;max-width:1100px}}svg{{display:block;width:100%;height:auto}}table{{border-collapse:collapse}}td,th{{border:1px solid #bbb;padding:.35rem .55rem;text-align:right}}td:first-child,th:first-child{{text-align:left}}code,pre{{background:#f4f4f4;padding:.2rem .35rem}}</style>
<h1>Least-squares refinement</h1>
<p><b>Mode:</b> {html.escape(mode)}</p>
<p><b>Reference chromatogram:</b> {html.escape(str(Path(chromatogram_csv).resolve()))}<br>
Detected x column: <code>{html.escape(ref.x_column)}</code> ({html.escape(ref.x_unit)}); signal column: <code>{html.escape(ref.signal_column)}</code>.</p>
{x_assumption_note}
{run_context_html}
<p><b>Optimiser:</b> SciPy bounded nonlinear least_squares, linear loss.<br>
Converged flag: {bool(opt.success)}; decision: {acceptance['status']}; applied to active model: {accepted}; function evaluations reported by optimiser: {int(opt.nfev)}; objective evaluations including final check: {eval_count}.</p>
<p><b>Acceptance checks:</b> {html.escape('; '.join(acceptance['reasons']) or 'All enabled checks passed.')} Bound-after-load minimum {config['least_squares']['minimum_bound_percent_after_load']:.3g}%; elution recovery minimum {config['least_squares']['minimum_elution_recovery_percent']:.3g}% of end-load column inventory. CPA policy: {html.escape(config['least_squares']['physical_acceptance_mode'])}.</p>
<p><b>Jacobian diagnostics:</b> {html.escape(json.dumps(jac_diagnostics, ensure_ascii=False))}</p>
<p><b>Initial total objective:</b> {initial_obj:.7g}<br><b>Final total objective:</b> {final_obj:.7g}<br>
<b>Improvement:</b> {improvement_percent:.4g}%<br><b>Final chromatogram objective ({html.escape(config['least_squares']['objective'])}):</b> {final_info['chromatogram_sse_normalized']:.7g}</p>
{comp_note}
<h2>Measured and predicted chromatogram</h2>
<p>The predicted UV and process overlays span the full programmed run. Measured UV is displayed as the median imported sample in each chart pixel column to prevent dark overdraw; every imported sample contributes to the chromatogram objective.</p>
{chromatogram_plot}
{fraction_uv_plot}
{composition_plot}
{_cpa_binding_html(binding_metrics)}
<h2>Refined mechanistic parameters</h2>
<table><thead><tr><th>Parameter</th><th>Initial</th><th>Fitted</th><th>Fit limit ±%</th><th>Allowed min</th><th>Allowed max</th></tr></thead><tbody>{parameter_rows}</tbody></table>
<h2>Objective construction</h2>
<p>The chromatogram residual is point-by-point least squares (Hahn et al., equation 9). RAW_SSE is the default. No baseline is inferred unless INITIAL_MEDIAN is explicitly selected; otherwise the fixed detector baseline is used. The predicted UV signal is the sum of each species concentration multiplied by its own response factor; the target uses the shared Model Parameters factor and each impurity uses its component-specific factor. UV references labelled AU are converted to mAU. Optional NORMALIZED_MSE divides residuals by the measured signal scale and sqrt(N). RAW_SSE compares UV values directly.</p>
{uv_identifiability_note}
<p>For each entered fraction, species mass% is 100 × its integrated detector concentration over the collected CV interval / the integrated concentration of all active species. CV is proportional to collected volume, so this accounts for changing flow. Legacy point-only references retain midpoint comparison. Each entered species value contributes one residual, divided by 100 and sqrt(the number of populated cells). With RAW_SSE, that residual group is converted to comparable UV units using the measured UV signal scale (at least 1 mAU) and sqrt(the number of UV samples); with NORMALIZED_MSE both groups stay dimensionless with a 1 mAU UV normalization floor. Blank cells add no residual and break the display line.</p>
<p>Only unlocked eligible mechanistic and instrument parameters listed above are refined. Locked values, column geometry and each recipe remain fixed. The candidate was saved to <code>{html.escape(fit_config_path.name)}</code>; it is applied to <code>{html.escape(CONFIG_PATH.name)}</code> only after all enabled acceptance checks pass.</p>
<h2>Objective improvement over time</h2>
<p>The plot and CSV show every objective evaluation, elapsed seconds, best-so-far objective, improvement relative to the initial objective, and any numerical-penalty status. Lower objective is better. <a href="{html.escape(fit_history_path.name)}">Open objective history CSV</a>.</p>
{history_plot}
<p>Files: <code>{fit_trace_path.name}</code>, <code>{fit_composition_path.name}</code>, <code>{fit_results_path.name}</code>, <code>{fit_history_path.name}</code>.</p>
""", encoding="utf-8")

    # Keep the existing "latest fit" files for the UI buttons while retaining
    # one permanent set per saved chromatogram profile.
    if run_context and fit_report_path != FIT_REPORT_HTML:
        for source_path, latest_path in (
            (fit_results_path, FIT_RESULTS_CSV),
            (fit_trace_path, FIT_TRACE_CSV),
            (fit_composition_path, FIT_COMPOSITION_CSV),
            (fit_report_path, FIT_REPORT_HTML),
            (fit_config_path, FIT_CONFIG_JSON),
            (fit_settings_path, FIT_SETTINGS_JSON),
            (fit_history_path, FIT_HISTORY_CSV),
        ):
            if source_path.is_file():
                shutil.copyfile(source_path, latest_path)
        if accepted and FIT_APPLY_JSL.is_file():
            shutil.copyfile(FIT_APPLY_JSL, fit_apply_profile_path)

    return {
        "success": bool(opt.success),
        "accepted": accepted,
        "acceptance": acceptance,
        "message": str(opt.message),
        "initial_objective": initial_obj,
        "final_objective": final_obj,
        "improvement_percent": improvement_percent,
        "chromatogram_sse": float(final_info["chromatogram_sse_normalized"]),
        "composition_sse": float(final_info["composition_sse_normalized"]),
        "nfev": int(opt.nfev),
        "residual_evaluations": eval_count,
        "mode": mode,
        "target_run_context": run_context,
        "fitted_config": fitted,
        "fit_results_csv": str(fit_results_path),
        "fit_trace_csv": str(fit_trace_path),
        "fit_composition_csv": str(fit_composition_path),
        "fit_report_html": str(fit_report_path),
        "fit_config_json": str(fit_config_path),
        "fit_settings_json": str(fit_settings_path),
        "fit_history_csv": str(fit_history_path),
        "latest_fit_history_csv": str(FIT_HISTORY_CSV),
        "apply_jsl": str(FIT_APPLY_JSL),
        "profile_apply_jsl": str(fit_apply_profile_path),
        "parameters": [s.path for s in specs],
    }


def run_global_least_squares_refinement(
    profiles: Iterable[dict[str, Any]],
    *,
    mode: str = MODE_CHROMATOGRAM,
    unlocked_paths: Iterable[str] | None = None,
    max_nfev: int = 60,
    max_model_evaluations: int | None = None,
) -> dict[str, Any]:
    """Fit one shared mechanistic parameter vector against selected run profiles.

    Each profile carries its own fixed recipe and CSV/Excel workbook. By default the point
    residuals are concatenated without scaling: their squared norm is equation
    9 summed across runs. NORMALIZED_MSE is an explicit optional weighting.
    """
    total_started = time.perf_counter()
    profiles = list(profiles)
    if not 1 <= len(profiles) <= MAX_LSQ_PROFILES:
        raise ValueError(f"Global least-squares fitting requires 1–{MAX_LSQ_PROFILES} attached chromatogram profiles.")
    mode = str(mode or "").strip().upper()
    if mode not in SUPPORTED_MODES:
        raise ValueError("Least-squares mode must be CHROMATOGRAM or CHROMATOGRAM_AND_COMPOSITION.")

    normalized_profiles: list[dict[str, Any]] = []
    candidate_paths: list[str] = []
    refs: list[ChromatogramReference] = []
    for index, profile in enumerate(profiles, start=1):
        cfg = normalize_config(profile.get("config"))
        if mode == MODE_CHROMATOGRAM_AND_COMPOSITION and cfg["least_squares"]["objective"] == "WEIGHTED_RMSE":
            raise ValueError("WEIGHTED_RMSE is defined for chromatogram signals only. Select RAW_SSE or NORMALIZED_MSE to include species mass% constraints.")
        errors = validate_config(cfg, strict=True)
        if errors:
            context = profile.get("run_context") or {}
            label = context.get("run_name") or f"Profile {index}"
            raise ValueError(f"Global-fit profile {label} has incomplete model inputs:\n- " + "\n- ".join(errors))
        source = profile.get("chromatogram_csv") or profile.get("csv_path")
        if not source:
            raise ValueError(f"Global-fit profile {index} has no chromatogram CSV path.")
        ref = load_chromatogram_reference(source, cfg,
                                         sheet_name=profile.get("sheet_name"),
                                         x_column=profile.get("x_column"),
                                         signal_column=profile.get("signal_column"),
                                         x_unit=profile.get("x_unit"),
                                         x_origin=profile.get("x_origin"),
                                         require_uv=True,
                                         fit_conductivity=bool(profile.get("fit_conductivity")),
                                         conductivity_column=profile.get("conductivity_column"))
        profile_weight = float(profile.get("weight", 1.0))
        if not math.isfinite(profile_weight) or profile_weight <= 0:
            raise ValueError(f"Global-fit profile {index}: run weight must be finite and > 0.")
        raw_composition_refs = profile.get("composition_references", [])
        profile_refs = validate_composition_references(cfg, raw_composition_refs) if mode == MODE_CHROMATOGRAM_AND_COMPOSITION and raw_composition_refs else []
        context = dict(profile.get("run_context") or {})
        context.setdefault("run_number", index)
        context.setdefault("run_name", f"Profile {index}")
        context.setdefault("model_run_source", "INDEPENDENT_LSQ")
        context.setdefault("model_run_number", context["run_number"])
        context.setdefault("model_run_name", context["run_name"])
        context.setdefault("model_run_label", f"Independent LSQ recipe Run {context['model_run_number']}")
        context["chromatogram_csv"] = str(Path(source).resolve())
        paths = [path for path in required_parameter_paths(cfg)
                 if path not in {"column.volume_mL", "column.length_mm"} and not path.endswith(".mass_percent")]
        for path in paths:
            if path not in candidate_paths:
                candidate_paths.append(path)
        normalized_profiles.append({
            "config": cfg,
            "reference": ref,
            "fit_conductivity": bool(profile.get("fit_conductivity")),
            "composition_references": profile_refs,
            "run_context": context,
            "weight": profile_weight,
            "x_origin": _validate_x_origin(profile.get("x_origin"), ref.x_unit),
            "sheet_name": ref.sheet_name or "CSV",
        })

    if mode == MODE_CHROMATOGRAM_AND_COMPOSITION and not any(p["composition_references"] for p in normalized_profiles):
        raise ValueError("Supply at least one composition reference for a selected profile when chromatogram + mass% is selected.")
    first = normalized_profiles[0]["config"]
    if mode == MODE_CHROMATOGRAM_AND_COMPOSITION and first["least_squares"]["objective"] == "WEIGHTED_RMSE":
        raise ValueError("WEIGHTED_RMSE is defined for chromatogram signals only. For additional species mass% constraints select RAW_SSE or NORMALIZED_MSE, because composition residuals have different units.")
    for index, profile in enumerate(normalized_profiles[1:], start=2):
        cfg = profile["config"]
        if cfg["model"] != first["model"] or int(cfg.get("impurity_count", 0)) != int(first.get("impurity_count", 0)):
            raise ValueError("All attached chromatograms must use the same model combination and selected species for one global shared fit.")
        if [r.get("name") for r in active_components(cfg)] != [r.get("name") for r in active_components(first)]:
            raise ValueError(f"Global-fit profile {index} does not have the same shared protein-species names as profile 1.")
        if cfg["least_squares"].get("fit_range_percent") != first["least_squares"].get("fit_range_percent"):
            raise ValueError(f"Global-fit profile {index} has different per-parameter fit ranges; one shared fit needs one shared set of limits.")
        acceptance_fields = ("physical_acceptance_mode", "minimum_bound_percent_after_load",
                             "minimum_elution_recovery_percent", "minimum_objective_improvement_percent",
                             "max_trial_minutes")
        if any(cfg["least_squares"].get(field) != first["least_squares"].get(field)
               for field in acceptance_fields):
            raise ValueError(f"Global-fit profile {index} has different physical acceptance settings.")

    # These are shared batch parameters. Process recipes and column geometry
    # remain profile inputs, while a different non-empty value for a shared
    # mechanistic parameter is rejected instead of silently choosing one run.
    seed = copy.deepcopy(first)
    for path in candidate_paths:
        values = []
        for profile in normalized_profiles:
            try:
                value = _get_path(profile["config"], path)
            except (KeyError, IndexError, TypeError):
                value = None
            if value is not None:
                values.append(value)
        if not values:
            continue
        reference_value = values[0]
        for value in values[1:]:
            try:
                same = math.isclose(float(reference_value), float(value), rel_tol=1e-10, abs_tol=1e-12)
            except (TypeError, ValueError):
                same = reference_value == value
            if not same:
                raise ValueError(f"Shared mechanistic parameter {path} differs between attached profiles. Global fitting requires one common starting value.")
        _set_path(seed, path, reference_value)
        for profile in normalized_profiles:
            _set_path(profile["config"], path, reference_value)

    if unlocked_paths is not None:
        unlocked_paths = list(unlocked_paths)
        for profile in normalized_profiles:
            profile["config"]["least_squares"]["unlocked_paths"] = list(unlocked_paths)
        seed["least_squares"]["unlocked_paths"] = list(unlocked_paths)
    # A later selected run may have a pH programme even when the first run is
    # isocratic. Use the recipe with the broadest active parameter contract to
    # build fit specs; all profiles already share the same mechanistic values.
    spec_config = max(normalized_profiles, key=lambda p: len(required_parameter_paths(p["config"])))["config"]
    specs = fit_parameter_specs(
        spec_config,
        fit_conductivity=any(p["fit_conductivity"] for p in normalized_profiles),
        unlocked_paths=unlocked_paths,
        candidate_paths=candidate_paths,
    )
    if max_model_evaluations is not None:
        if int(max_model_evaluations) != max_model_evaluations or max_model_evaluations < 2:
            raise ValueError("Model-evaluation safety cap must be a whole number of at least 2.")
        # One candidate and a complete numerical Jacobian are needed before
        # SciPy can make a meaningful parameter step. Keep one evaluation for
        # final verification, which is included in the displayed count.
        required = len(specs) + 2 if specs else 1
        if max_model_evaluations < required:
            raise ValueError(
                f"Model-evaluation safety cap {max_model_evaluations} is too small for "
                f"{len(specs)} unlocked parameters. Select at least {required} to allow "
                "one complete Jacobian and final verification."
            )
    x0, lo, hi = _encode_initial(specs)
    n_profiles = len(normalized_profiles)
    normalized_objective = first["least_squares"]["objective"] == "NORMALIZED_MSE"
    weighted_objective = first["least_squares"]["objective"] == "WEIGHTED_RMSE"
    run_weights = np.array([p["weight"] for p in normalized_profiles], dtype=float)
    run_weights /= float(np.sum(run_weights))
    for profile in normalized_profiles:
        if profile["config"]["least_squares"]["objective"] != first["least_squares"]["objective"] or profile["config"]["chromatography_mode"] != first["chromatography_mode"]:
            raise ValueError("Global profiles must use the same objective and chromatography mode.")
    def evaluate(configs: list[dict[str, Any]], *, detailed: bool = True) -> tuple[np.ndarray, list[dict[str, Any] | None]]:
        residual_blocks = []
        infos = []
        for cfg, profile, normalized_weight in zip(configs, normalized_profiles, run_weights):
            deadline = _trial_deadline(cfg)
            r, info = _objective_parts(cfg, profile["reference"], mode,
                                       profile["composition_references"], profile["fit_conductivity"],
                                       **({"solver_deadline": deadline} if deadline is not None else {}))
            # Preserve the original objective and run weighting exactly.
            scale_factor = math.sqrt(float(normalized_weight)) if weighted_objective else (1.0 / math.sqrt(n_profiles) if normalized_objective else 1.0)
            residual_blocks.append(r * scale_factor)
            infos.append(info if detailed else None)
        return np.concatenate(residual_blocks), infos

    base_configs = [copy.deepcopy(p["config"]) for p in normalized_profiles]
    active_path = FIT_HISTORY_CSV.with_name("least_squares_active_evaluation.txt")
    try:
        active_path.write_text(
            f"Initial model evaluation started {time.strftime('%Y-%m-%d %H:%M:%S %Z')}\n"
            f"TDM+CPA trial timeout: {first['least_squares']['max_trial_minutes']:g} minutes (0 = off)\n",
            encoding="utf-8")
    except OSError:
        pass
    initial_r, initial_infos = evaluate(base_configs)
    initial_obj = float(np.sum(initial_r ** 2))
    fit_started = time.perf_counter()
    best_objective = initial_obj
    best_vector = None  # No improved candidate: retain exact original numeric values.
    history: list[dict[str, Any]] = [{
        "evaluation": 0, "elapsed_seconds": 0.0, "objective": initial_obj,
        "best_objective": initial_obj, "improvement_percent": 0.0,
        "phase": "initial", "status": "ok",
    }]
    eval_count = 0
    candidate_evaluations = 0
    budget_reached = False
    invalid_trial_vectors: set[bytes] = set()
    timed_out_trial_vectors: set[bytes] = set()
    optimizer_budget = None if max_model_evaluations is None else int(max_model_evaluations) - 1
    history_fields = ["evaluation", "elapsed_seconds", "objective", "best_objective", "improvement_percent", "phase", "status"]
    history_path: Path | None = FIT_HISTORY_CSV
    current_probe_label = ""
    history_backup = FIT_HISTORY_CSV.with_name(
        f"{FIT_HISTORY_CSV.stem}_recovered_{os.getpid()}_{time.time_ns()}.csv")
    history_temp = Path(tempfile.gettempdir()) / history_backup.name
    history_logging_disabled = False
    history_write_error: str | None = None

    def write_active(phase: str) -> None:
        label = current_probe_label if phase == "jacobian" else ""
        try:
            active_path.write_text(
                f"Evaluation {eval_count} started {time.strftime('%Y-%m-%d %H:%M:%S %Z')}\n"
                f"Phase: {phase}\n"
                + (f"Parameter: {label}\n" if label else "")
                + f"TDM+CPA trial timeout: {first['least_squares']['max_trial_minutes']:g} minutes (0 = off)\n",
                encoding="utf-8")
        except OSError:
            pass  # Progress I/O cannot invalidate a solver evaluation.

    def write_history() -> None:
        nonlocal history_path, history_logging_disabled, history_write_error
        if history_logging_disabled:
            return
        targets = dict.fromkeys([p for p in (history_path, history_backup, history_temp) if p is not None])
        for path in targets:
            try:
                path.parent.mkdir(parents=True, exist_ok=True)
                with path.open("w", encoding="utf-8", newline="") as fh:
                    writer = csv.DictWriter(fh, fieldnames=history_fields)
                    writer.writeheader()
                    writer.writerows(history)
            except OSError as exc:
                history_write_error = f"{type(exc).__name__}: {exc}"
                continue
            history_path = path
            return
        history_path = None
        history_logging_disabled = True

    def append_history(row: dict[str, Any]) -> None:
        nonlocal history_path, history_write_error
        if history_path is None:
            return
        try:
            with history_path.open("a", encoding="utf-8", newline="") as fh:
                csv.DictWriter(fh, fieldnames=history_fields).writerow(row)
        except OSError as exc:
            # Progress I/O is separate from the scientific objective. A locked
            # OneDrive/Excel CSV must not turn a valid trial into a penalty.
            history_write_error = f"{type(exc).__name__}: {exc}"
            history_path = history_backup if history_path == FIT_HISTORY_CSV else history_temp
            write_history()  # Restore every row, including the failed append.

    def record(objective: float, phase: str, status: str) -> None:
        nonlocal best_objective
        best_objective = min(best_objective, objective)
        row = {
            "evaluation": eval_count,
            "elapsed_seconds": time.perf_counter() - fit_started,
            "objective": objective,
            "best_objective": best_objective,
            "improvement_percent": _objective_improvement_percent(initial_obj, best_objective),
            "phase": phase,
            "status": status,
        }
        history.append(row)
        # Normal progress writes append one row; only a locked file triggers
        # a full in-memory snapshot to a fresh path.
        append_history(row)

    write_history()

    def residual(vector: np.ndarray, phase: str = "optimizer") -> np.ndarray:
        nonlocal eval_count, best_vector, candidate_evaluations
        if optimizer_budget is not None and eval_count >= optimizer_budget:
            raise EvaluationBudgetReached("Selected model-evaluation safety cap reached.")
        eval_count += 1
        if phase == "optimizer":
            candidate_evaluations += 1
        write_active(phase)
        try:
            candidates = [_apply_vector(cfg, specs, vector) for cfg in base_configs]
            values, _infos = evaluate(candidates, detailed=False)
            objective = float(np.sum(values * values))
            if not np.all(np.isfinite(values)) or not math.isfinite(objective):
                raise FloatingPointError("non-finite global residual")
            invalid_trial_vectors.discard(np.asarray(vector, dtype=float).tobytes())
            timed_out_trial_vectors.discard(np.asarray(vector, dtype=float).tobytes())
            if objective < best_objective:
                best_vector = np.asarray(vector).copy()
            record(objective, phase, "ok")
            return values
        except EvaluationBudgetReached:
            raise
        except Exception as exc:
            invalid_trial_vectors.add(np.asarray(vector, dtype=float).tobytes())
            if isinstance(exc, TimeoutError):
                timed_out_trial_vectors.add(np.asarray(vector, dtype=float).tobytes())
            penalty = np.full_like(initial_r, 1e6, dtype=float)
            record(float(np.sum(penalty * penalty)), phase, f"penalty: {type(exc).__name__}: {exc}")
            return penalty

    def serial_jacobian(vectors: list[np.ndarray]) -> list[np.ndarray]:
        nonlocal current_probe_label
        results = []
        for index, vector in enumerate(vectors):
            if len(vectors) == len(specs):
                current_probe_label = specs[index].label
            results.append(residual(vector, "jacobian"))
        return results

    def label_probe(index: int) -> None:
        nonlocal current_probe_label
        current_probe_label = specs[index].label

    if specs:
        cached_residual = _cache_last_evaluation(residual)
        try:
            opt = least_squares(
                cached_residual,
                x0,
                bounds=(lo, hi),
                jac=_bounded_jacobian(
                    cached_residual, lo, hi, serial_jacobian,
                    trial_valid=lambda vector: np.asarray(vector, dtype=float).tobytes()
                    not in invalid_trial_vectors,
                    trial_timed_out=lambda vector: np.asarray(vector, dtype=float).tobytes()
                    in timed_out_trial_vectors,
                    parameter_labels=[spec.label for spec in specs],
                    probe_callback=label_probe),
                method=_optimizer_method(x0, lo, hi),
                loss="linear",
                x_scale="jac",
                ftol=1e-4,
                xtol=1e-4,
                gtol=1e-7,
                max_nfev=max(2, int(max_nfev)),
            )
        except EvaluationBudgetReached as exc:
            budget_reached = True
            opt = SimpleNamespace(x=x0.copy(), success=False, message=str(exc),
                                  nfev=candidate_evaluations)
        except TimeoutError as exc:
            try:
                active_path.write_text(
                    f"Least-squares fit stopped: {exc}\n"
                    f"Completed model residual evaluations: {eval_count}.\n",
                    encoding="utf-8")
            except OSError:
                pass
            raise
        jac_diagnostics = _jacobian_diagnostics(opt, best_vector if best_vector is not None else x0, specs)
        if best_vector is None:
            fitted_configs = [copy.deepcopy(cfg) for cfg in base_configs]
            opt.x = np.asarray([math.log10(float(_get_path(seed, spec.path))) if spec.transform == "LOG10"
                                else float(_get_path(seed, spec.path)) for spec in specs])
        else:
            opt.x = best_vector.copy()
            fitted_configs = [_apply_vector(cfg, specs, opt.x) for cfg in base_configs]
        opt_success = bool(opt.success)
        opt_message = str(opt.message)
        opt_nfev = int(opt.nfev)
    else:
        fitted_configs = [copy.deepcopy(cfg) for cfg in base_configs]
        opt_success = True
        opt_message = "All eligible mechanistic parameters are locked; objective evaluated without optimization."
        opt_nfev = 0
        jac_diagnostics = _jacobian_diagnostics(SimpleNamespace(), x0, specs)
    try:
        final_r, final_infos = evaluate(fitted_configs)
        final_obj = float(np.sum(final_r ** 2))
        if not math.isfinite(final_obj):
            raise FloatingPointError("non-finite final global objective")
    except Exception as exc:
        fitted_configs = [copy.deepcopy(cfg) for cfg in base_configs]
        final_r, final_infos = initial_r, initial_infos
        final_obj = initial_obj
        opt_success = False
        opt_message = f"Final candidate verification failed; original parameters retained: {type(exc).__name__}: {exc}"
    eval_count += 1
    record(final_obj, "final verification", "ok" if opt_success else opt_message)
    acceptance = _fit_acceptance(first, initial_obj, final_obj, final_infos)
    accepted = acceptance["accepted"]
    fitted = fitted_configs[0]

    save_config(fitted, FIT_CONFIG_JSON)
    if accepted:
        save_config(fitted, CONFIG_PATH)
        _write_apply_jsl(fitted, specs)
    elif FIT_APPLY_JSL.exists():
        FIT_APPLY_JSL.unlink()

    with FIT_RESULTS_CSV.open("w", encoding="utf-8", newline="") as fh:
        fields = ["Parameter_Path", "Parameter", "Initial_Value", "Fitted_Value", "Transform", "Fit_State",
                  "Fit_Range_Percent", "Allowed_Min", "Allowed_Max", "At_Lower_Bound", "At_Upper_Bound"]
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        spec_map = {spec.path: spec for spec in specs}
        for path in candidate_paths:
            try:
                initial_value = float(_get_path(seed, path))
                fitted_value = float(_get_path(fitted, path))
            except (KeyError, TypeError, ValueError):
                continue
            spec = spec_map.get(path)
            fit_state = "Unlocked" if spec else "Locked"
            xfit = None
            if spec and specs:
                xfit = (10.0 ** float(opt.x[specs.index(spec)]) if spec.transform == "LOG10"
                        else float(opt.x[specs.index(spec)]))
            allowed_min, allowed_max = _physical_fit_bounds(spec) if spec else ("", "")
            writer.writerow({
                "Parameter_Path": path, "Parameter": spec.label if spec else _report_parameter_label(seed, path),
                "Initial_Value": initial_value, "Fitted_Value": fitted_value,
                "Transform": spec.transform if spec else "", "Fit_State": fit_state,
                "Fit_Range_Percent": first["least_squares"].get("fit_range_percent", {}).get(path, "") if spec else "",
                "Allowed_Min": allowed_min, "Allowed_Max": allowed_max,
                "At_Lower_Bound": bool(spec and xfit is not None and abs(float(opt.x[specs.index(spec)]) - spec.lower) <= 1e-6 * max(1.0, abs(spec.lower))),
                "At_Upper_Bound": bool(spec and xfit is not None and abs(float(opt.x[specs.index(spec)]) - spec.upper) <= 1e-6 * max(1.0, abs(spec.upper))),
            })

    # Only unlocked parameters are listed as improvements; report the exact
    # values in original physical units, even when the fit variable was log10.
    with FIT_IMPROVED_CSV.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=["Parameter_Path", "Parameter", "Initial_Value", "Improved_Value", "Change_percent",
                                                "Fit_Range_Percent", "Allowed_Min", "Allowed_Max", "At_Lower_Bound", "At_Upper_Bound"])
        writer.writeheader()
        text_rows = ["LEAST-SQUARES REFINED PARAMETERS", f"Objective: {first['least_squares']['objective']}", f"Accepted: {accepted}", f"Initial objective: {initial_obj:.10g}", f"Final objective: {final_obj:.10g}", ""]
        for i, spec in enumerate(specs):
            initial_val = float(_get_path(seed, spec.path))
            fitted_val = float(_get_path(fitted, spec.path))
            diff_percent = 100.0 * (fitted_val - initial_val) / abs(initial_val) if initial_val != 0 else math.nan
            xfit_val = math.log10(fitted_val) if spec.transform == "LOG10" else fitted_val
            row = {"Parameter_Path": spec.path, "Parameter": spec.label,
                   "Initial_Value": initial_val, "Improved_Value": fitted_val,
                   "Change_percent": diff_percent,
                   "Fit_Range_Percent": first["least_squares"].get("fit_range_percent", {}).get(spec.path, ""),
                   "Allowed_Min": _physical_fit_bounds(spec)[0], "Allowed_Max": _physical_fit_bounds(spec)[1],
                   "At_Lower_Bound": abs(xfit_val - spec.lower) <= 1e-6 * max(1.0, abs(spec.lower)),
                   "At_Upper_Bound": abs(xfit_val - spec.upper) <= 1e-6 * max(1.0, abs(spec.upper))}
            writer.writerow(row)
            text_rows.append(f"{spec.path} ({spec.label}): {initial_val:.8g} -> {fitted_val:.8g} ({diff_percent:+.5g}%)")
        if not specs:
            text_rows.append("No parameters unlocked for refinement.")
        FIT_IMPROVED_TXT.write_text("\n".join(text_rows) + "\n", encoding="utf-8")

    metric_rows: list[dict[str, Any]] = []
    with FIT_TRACE_CSV.open("w", encoding="utf-8", newline="") as fh:
        trace_fields = ["Run_Number", "Run_Name", "Model_Run", "Model_Run_Name", "Chromatogram_CSV", "CV", "Measured_Signal",
                        "Simulated_Total_Protein_g_L", "Predicted_Detector_Signal", "Residual",
                        "Normalized_Residual", "Sample_Weight", "Run_Weight", "Profile_Objective_Contribution"]
        writer = csv.DictWriter(fh, fieldnames=trace_fields)
        writer.writeheader()
        for profile, initial_info, info, norm_run_weight in zip(normalized_profiles, initial_infos, final_infos, run_weights):
            chrom = info["chromatogram"]
            context = profile["run_context"]
            difference = np.asarray(chrom["predicted_detector_signal"], dtype=float) - np.asarray(chrom["measured"], dtype=float)
            norm_difference = difference / float(chrom["measured_signal_scale"])
            normalized_nrmse = float(np.sqrt(np.mean(difference ** 2))) / float(chrom["measured_signal_scale"])
            raw_rmse = float(np.sqrt(np.mean(difference ** 2)))
            initial_chrom = initial_info["chromatogram"]
            initial_difference = np.asarray(initial_chrom["predicted_detector_signal"], dtype=float) - np.asarray(initial_chrom["measured"], dtype=float)
            initial_raw_rmse = float(np.sqrt(np.mean(initial_difference ** 2)))
            sample_weights = np.asarray(chrom["point_weights"], dtype=float)
            run_weighted_rmse = float(np.sqrt(np.sum(sample_weights * difference ** 2) / np.sum(sample_weights)))
            total_ss = float(np.sum((np.asarray(chrom["measured"], dtype=float) - float(np.mean(chrom["measured"]))) ** 2))
            r2 = 1.0 - float(np.sum(difference ** 2)) / total_ss if total_ss > 1e-30 else math.nan
            metric = {
                "run_number": int(context["run_number"]),
                "run_name": str(context["run_name"]),
                "model_run_label": str(context["model_run_label"]),
                "model_run_name": str(context["model_run_name"]),
                "chromatogram_csv": context["chromatogram_csv"],
                "overlap_points": int(chrom["overlap_points"]),
                "normalized_nrmse": normalized_nrmse,
                "raw_rmse": raw_rmse,
                "initial_raw_rmse": initial_raw_rmse,
                "weighted_rmse": run_weighted_rmse,
                "run_weight": profile["weight"],
                "r_squared": r2,
                "conductivity_rmse_mS_cm": info["conductivity"]["rmse_mS_cm"] if info["conductivity"] else None,
                "initial_uv_objective": float(initial_info["chromatogram_sse_normalized"]),
                "fitted_uv_objective": float(info["chromatogram_sse_normalized"]),
                "initial_mass_objective": float(initial_info["composition_sse_normalized"]),
                "fitted_mass_objective": float(info["composition_sse_normalized"]),
                "profile_objective": float(info["chromatogram_sse_normalized"] + info["conductivity_sse_normalized"] + info["composition_sse_normalized"]),
            }
            metric_rows.append(metric)
            for cv, measured, sim_total, predicted, residual_value, norm_value, sample_weight in zip(
                chrom["CV"], chrom["measured"], chrom["simulated_total_g_L"],
                chrom["predicted_detector_signal"], difference, norm_difference, sample_weights,
            ):
                writer.writerow({
                    "Run_Number": metric["run_number"], "Run_Name": metric["run_name"],
                    "Model_Run": metric["model_run_label"], "Model_Run_Name": metric["model_run_name"],
                    "Chromatogram_CSV": metric["chromatogram_csv"], "CV": float(cv),
                    "Measured_Signal": float(measured), "Simulated_Total_Protein_g_L": float(sim_total),
                    "Predicted_Detector_Signal": float(predicted), "Residual": float(residual_value),
                    "Normalized_Residual": float(norm_value),
                    "Sample_Weight": float(sample_weight), "Run_Weight": float(profile["weight"]),
                    "Profile_Objective_Contribution": (float(norm_run_weight * sample_weight / np.sum(sample_weights) * residual_value ** 2)
                         if weighted_objective else float(norm_value ** 2 / len(chrom["CV"]) / n_profiles) if normalized_objective else float(residual_value ** 2)),
                })

    composition_rows = []
    composition_profile_metrics: list[dict[str, Any]] = []
    for profile, initial_info, final_info in zip(normalized_profiles, initial_infos, final_infos):
        initial_by_fraction = {
            (row["Fraction"], row["Protein_Species"]): row
            for row in initial_info["composition"]
        }
        for species in active_components(profile["config"]):
            name = str(species["name"])
            final_species = [row for row in final_info["composition"]
                             if row["Protein_Species"] == name]
            if not final_species:
                continue
            before = np.asarray([
                float(row["Measured_mass_percent"])
                - float(initial_by_fraction[(row["Fraction"], name)]["Simulated_mass_percent"])
                for row in final_species], dtype=float)
            after = np.asarray([
                float(row["Measured_mass_percent"])
                - float(row["Simulated_mass_percent"])
                for row in final_species], dtype=float)
            composition_profile_metrics.append({
                "run_number": int(profile["run_context"]["run_number"]),
                "run_name": str(profile["run_context"]["run_name"]),
                "protein_species": name,
                "observed_fractions": len(final_species),
                "initial_rmse_percentage_points": float(np.sqrt(np.mean(before ** 2))),
                "final_rmse_percentage_points": float(np.sqrt(np.mean(after ** 2))),
            })
    conductivity_output_path = FIT_TRACE_CSV.with_name(FIT_CONDUCTIVITY_CSV.name)
    with conductivity_output_path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=["Run_Number", "Run_Name", "Model_Run", "Conductivity_Column", "CV", "Measured_mS_cm", "Predicted_mS_cm", "Residual_mS_cm"])
        writer.writeheader()
        for profile, info in zip(normalized_profiles, final_infos):
            cond = info["conductivity"]
            if not cond:
                continue
            context = profile["run_context"]
            for cv, measured, predicted in zip(cond["CV"], cond["measured"], cond["predicted"]):
                writer.writerow({"Run_Number": context["run_number"], "Run_Name": context["run_name"],
                                 "Model_Run": context["model_run_label"], "Conductivity_Column": cond["column"],
                                 "CV": float(cv), "Measured_mS_cm": float(measured),
                                 "Predicted_mS_cm": float(predicted), "Residual_mS_cm": float(predicted - measured)})
    with FIT_COMPOSITION_CSV.open("w", encoding="utf-8", newline="") as fh:
        fields = ["Run_Number", "Run_Name", "CV", "Fraction", "Fraction_Start", "Fraction_End", "Fraction_Unit", "Protein_Species", "Measured_mass_percent", "Simulated_mass_percent", "Model_Fraction_Quantifiable", "Residual_mass_percent"]
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        for profile, info in zip(normalized_profiles, final_infos):
            context = profile["run_context"]
            for raw in info["composition"]:
                row = {
                    "Run_Number": context["run_number"], "Run_Name": context["run_name"],
                    **raw,
                    "Residual_mass_percent": float(raw["Measured_mass_percent"]) - float(raw["Simulated_mass_percent"]),
                }
                composition_rows.append(row)
                writer.writerow(row)

    write_history()
    binding_metrics_by_run = [
        {"run_number": int(profile["run_context"]["run_number"]),
         "run_name": str(profile["run_context"]["run_name"]),
         "audit": _cpa_binding_metrics(initial_info, final_info)}
        for profile, initial_info, final_info in zip(normalized_profiles, initial_infos, final_infos)
    ]
    settings = {
        "mode": mode,
        "profile_count": n_profiles,
        "objective": first["least_squares"]["objective"],
        "equation_9": "min_theta sum_runs sum_UV_points (predicted_UV - measured_UV)^2; enabled conductivity and mass% add separately scaled residual groups",
        "weighted_rmse_formula": "sqrt(sum_runs (run_weight/sum_run_weights) * sum_points(point_weight*residual^2)/sum_point_weights)",
        "optimizer": f"scipy.optimize.least_squares(method={_optimizer_method(x0, lo, hi)}, loss=linear, ftol=1e-4, xtol=1e-4, gtol=1e-7); bounded residual-vector least squares",
        "baseline_mode": first["least_squares"]["baseline_mode"],
        "profile_weighting": "User run weights times within-run weighted MSE; sqrt gives weighted RMSE" if weighted_objective else "Equal normalized profile MSE" if normalized_objective else "Raw UV point SSE (equation 9), plus separately normalized conductivity and mass% when enabled",
        "profiles": [
            {
                "run_context": p["run_context"],
                "x_column": p["reference"].x_column,
                "signal_column": p["reference"].signal_column,
                "fit_conductivity": p["fit_conductivity"],
                "conductivity_column": p["reference"].conductivity_column,
                "x_unit": p["reference"].x_unit,
                "x_origin": p["x_origin"],
                "x_axis_assumption": p["reference"].x_axis_assumption,
                "selected_sheet": p["reference"].sheet_name or "CSV",
                "composition_references": p["composition_references"],
                "run_weight": p["weight"], "point_weight_column": p["reference"].weight_column,
                "recipe": p["config"],
                "baseline_and_objective": p["config"]["least_squares"],
            }
            for p in normalized_profiles
        ],
        "fitted_parameter_paths": [s.path for s in specs],
        "fit_range_percent": first["least_squares"].get("fit_range_percent", {}),
        "locked_parameter_paths": [p for p in candidate_paths if p not in {s.path for s in specs}],
        "max_nfev": int(max_nfev),
        "max_model_evaluations": int(max_model_evaluations) if max_model_evaluations is not None else None,
        "max_trial_minutes": float(first["least_squares"]["max_trial_minutes"]),
        "evaluation_cap_reached": budget_reached,
        "composition_profile_metrics": composition_profile_metrics,
        "CPA_binding_audit_by_run": binding_metrics_by_run,
        "acceptance": acceptance,
        "jacobian_diagnostics": jac_diagnostics,
        "objective_history_csv": str(history_path) if history_path else None,
        "objective_history_write_error": history_write_error,
    }
    parameter_rows = "".join(
        f"<tr><td>{html.escape(spec.label)}</td><td>{float(_get_path(seed, spec.path)):.7g}</td><td>{float(_get_path(fitted, spec.path)):.7g}</td>"
        f"<td>{html.escape(str(first['least_squares'].get('fit_range_percent', {}).get(spec.path, 'Default')))}</td>"
        f"<td>{_physical_fit_bounds(spec)[0]:.7g}</td><td>{_physical_fit_bounds(spec)[1]:.7g}</td></tr>"
        for spec in specs
    ) or '<tr><td colspan="4">No parameters were unlocked; model values were evaluated without optimization.</td></tr>'
    profile_rows = "".join(
        "<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in (
            metric["run_number"], metric["run_name"], metric["model_run_label"], metric["overlap_points"],
            f'{metric["normalized_nrmse"]:.6g}', f'{metric["initial_raw_rmse"]:.6g}', f'{metric["raw_rmse"]:.6g}', f'{metric["weighted_rmse"]:.6g}', f'{metric["run_weight"]:.6g}',
            "—" if not math.isfinite(metric["r_squared"]) else f'{metric["r_squared"]:.6g}',
            "—" if metric["conductivity_rmse_mS_cm"] is None else f'{metric["conductivity_rmse_mS_cm"]:.6g}',
        )) + "</tr>"
        for metric in metric_rows
    )
    composition_metric_rows = "".join(
        "<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in (
            metric["run_number"], metric["run_name"], metric["protein_species"],
            metric["observed_fractions"],
            f'{metric["initial_rmse_percentage_points"]:.6g}',
            f'{metric["final_rmse_percentage_points"]:.6g}',
        )) + "</tr>"
        for metric in composition_profile_metrics
    )
    channel_rows = "".join(
        "<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in (
            metric["run_number"], f'{metric["initial_uv_objective"]:.6g}',
            f'{metric["fitted_uv_objective"]:.6g}',
            f'{metric["initial_mass_objective"]:.6g}',
            f'{metric["fitted_mass_objective"]:.6g}',
        )) + "</tr>" for metric in metric_rows
    )
    composition_metrics_html = (
        "<h2>Fraction mass% fit by protein species</h2>"
        "<p>RMSE is measured in mass percentage points across collected fraction intervals. "
        "Blank species cells and fraction gaps add no observations.</p>"
        "<table><thead><tr><th>Run</th><th>Profile</th><th>Protein species</th>"
        "<th>Observed fractions</th><th>Before RMSE</th><th>After RMSE</th>"
        "</tr></thead><tbody>" + composition_metric_rows + "</tbody></table>"
    ) if composition_profile_metrics else ""
    recipe_details = "".join(
        f'<details><summary>Run {p["run_context"]["run_number"]} — {html.escape(p["run_context"]["run_name"])} recipe and input CSV</summary>'
        f'<pre>{html.escape(json.dumps({"run_context": p["run_context"], "configuration": p["config"]}, indent=2, ensure_ascii=False))}</pre></details>'
        for p in normalized_profiles
    )
    profile_plots = []
    for profile, cfg, info, binding_row in zip(normalized_profiles, fitted_configs, final_infos,
                                                binding_metrics_by_run):
        plot = _chromatogram_fit_svg(cfg, info)
        if info["conductivity"]:
            plot += _conductivity_fit_svg(info)
        refs = profile["composition_references"]
        if refs:
            plot += ("<h4>UV chromatogram with collected fraction mass%</h4>"
                     + _fraction_uv_overlay_svg(cfg, refs, info)
                     + "<h4>Fraction mass% at midpoints</h4>"
                     + _composition_profile_svg(cfg, refs, info["result"]["trace"])
                     + "<p>Only populated species cells for collected fraction intervals contribute mass% residuals. The measured target is solid, measured impurities are dotted, model target is dashed, and model impurities use lighter dots. Blank cells leave gaps; connecting lines are for display.</p>")
        profile_plots.append(
            f'<h3>Run {profile["run_context"]["run_number"]}: {html.escape(profile["run_context"]["run_name"])}'
            f' vs {html.escape(profile["run_context"]["model_run_label"])}</h3>'
            + plot + _cpa_binding_html(binding_row["audit"]))
    overlays_html = "".join(profile_plots)
    improvement = _objective_improvement_percent(initial_obj, final_obj)
    fit_quality_warning = ""
    poor_uv_runs = [str(row["run_number"]) for row in metric_rows
                    if math.isfinite(row["r_squared"]) and row["r_squared"] < 0.8]
    worsened_uv_runs = [str(row["run_number"]) for row in metric_rows
                        if row["raw_rmse"] > row["initial_raw_rmse"] * 1.05 + 1e-9]
    if poor_uv_runs:
        fit_quality_warning = ("Measured and fitted UV remain materially different for run(s) "
                               + ", ".join(poor_uv_runs)
                               + " (UV R² below 0.8). Acceptance checks do not certify a close UV match; "
                               "it does not mean the measured chromatogram was matched.")
    elif worsened_uv_runs:
        fit_quality_warning = ("Fitted UV RMSE is more than 5% worse than its starting value for run(s) "
                               + ", ".join(worsened_uv_runs)
                               + ". The mass% and UV constraints may conflict; review the plotted traces and fixed inputs.")
    elif budget_reached:
        fit_quality_warning = ("The model-evaluation safety cap stopped optimization before convergence. "
                               "The best tested parameters were retained.")
    binding_runs = [str(row["run_number"]) for row in binding_metrics_by_run
                    if row["audit"] and (row["audit"]["target_peak_before_elution"]
                                         or row["audit"]["fitted"]["species"][0]
                                         ["end_load_bound_percent_of_load"] < 1.0)]
    if binding_runs:
        fit_quality_warning += (" " if fit_quality_warning else "") + (
            "TDM+CPA run(s) " + ", ".join(binding_runs)
            + " do not predict convincing target binding and elution: the target outlet peak "
            "is before elution or less than 1% of its load is bound after loading. "
            "Review the CPA binding audit before treating the fitted parameters as physical estimates."
        )
    history_plot = _objective_history_svg(history)
    if history_path is None:
        history_file_html = "History CSV could not be written; the complete history is plotted below."
    elif history_path.parent == FIT_REPORT_HTML.parent:
        history_name = html.escape(history_path.name, quote=True)
        history_file_html = f'History file: <a href="{history_name}">{history_name}</a>'
    else:
        history_file_html = f"History file: <code>{html.escape(str(history_path))}</code>"
    if history_write_error:
        history_file_html += (". A progress-file write failed: "
                              + html.escape(history_write_error))
    fit_elapsed_seconds = time.perf_counter() - total_started
    fit_elapsed_hms = _format_elapsed_hms(fit_elapsed_seconds)
    settings["fit_elapsed_seconds"] = fit_elapsed_seconds
    settings["fit_elapsed_hms"] = fit_elapsed_hms
    settings["model_residual_evaluations"] = eval_count
    settings["fit_quality_warning"] = fit_quality_warning
    try:
        active_path.write_text(
            f"Least-squares fit completed {time.strftime('%Y-%m-%d %H:%M:%S %Z')}.\n"
            f"Model residual evaluations: {eval_count}.\n", encoding="utf-8")
    except OSError:
        pass
    FIT_SETTINGS_JSON.write_text(json.dumps(settings, indent=2, ensure_ascii=False), encoding="utf-8")
    FIT_REPORT_HTML.write_text(f"""<!doctype html>
<meta charset="utf-8"><title>Global least-squares refinement</title>
<style>body{{font-family:Arial,sans-serif;margin:2rem;max-width:1100px}}svg{{display:block;width:100%;height:auto}}table{{border-collapse:collapse}}td,th{{border:1px solid #bbb;padding:.35rem .55rem;text-align:right}}td:first-child,th:first-child{{text-align:left}}code,pre{{background:#f4f4f4;padding:.2rem .35rem}}details{{margin:.6rem 0}}</style>
<h1>Global least-squares refinement</h1>
<p><b>Elapsed fit time:</b> {fit_elapsed_hms} ({fit_elapsed_seconds:.1f} seconds). Measured from loading the selected reference profiles through fitted output preparation.</p>
<p><b>Mode:</b> {html.escape(mode)}; <b>Profiles:</b> {n_profiles}; <b>Model:</b> {html.escape(first["model"])}</p>
<p><b>Initial global objective:</b> {initial_obj:.7g}<br><b>Final global objective:</b> {final_obj:.7g}<br><b>Improvement:</b> {improvement:.4g}%<br>
<b>Decision:</b> {acceptance['status']}; <b>Applied:</b> {accepted}; <b>Optimizer converged:</b> {opt_success}; {html.escape(opt_message)}<br><b>Model residual evaluations:</b> {eval_count}; <b>Optimizer candidate evaluations:</b> {opt_nfev}; <b>Evaluation safety cap:</b> {int(max_model_evaluations) if max_model_evaluations is not None else int(max_nfev)}; <b>Unlocked parameters:</b> {len(specs)}. The model-evaluation cap includes optimizer candidates, finite-difference Jacobian probes, and final verification; the initial reference calculation is logged at evaluation 0.</p>
<p><b>Acceptance checks:</b> {html.escape('; '.join(acceptance['reasons']) or 'All enabled checks passed.')} Bound-after-load minimum {first['least_squares']['minimum_bound_percent_after_load']:.3g}%; elution recovery minimum {first['least_squares']['minimum_elution_recovery_percent']:.3g}% of end-load column inventory. CPA policy: {html.escape(first['least_squares']['physical_acceptance_mode'])}. Candidate parameters are saved for diagnosis and applied to the active model only after all enabled checks pass.</p>
<p><b>Local Jacobian diagnostics:</b> {html.escape(json.dumps(jac_diagnostics, ensure_ascii=False))}</p>
{f'<p><b>Fit quality warning:</b> {html.escape(fit_quality_warning)}</p>' if fit_quality_warning else ''}
<h2>Per-profile chromatogram metrics</h2><p>Objective: {html.escape(first["least_squares"]["objective"])}. Baseline policy: {html.escape(first["least_squares"]["baseline_mode"])}. Default RAW_SSE uses the exact pointwise UV equation 9; optional NORMALIZED_MSE gives equal normalized profile weights. UV is predicted as the sum of each species concentration multiplied by its own response factor. Conductivity, when enabled for a run, adds measured detector residuals normalized by that run's conductivity signal scale and sample count; in RAW_SSE these are rescaled to the measured UV scale.</p>
<table><thead><tr><th>Run</th><th>Profile</th><th>Model run</th><th>Overlap points</th><th>Normalized NRMSE</th><th>Initial UV RMSE</th><th>Fitted UV RMSE</th><th>Weighted UV RMSE</th><th>Run weight</th><th>UV R²</th><th>Conductivity RMSE [mS/cm]</th></tr></thead><tbody>{profile_rows}</tbody></table>
<h2>Objective by measurement type</h2><p>Per-run squared residual sums before optional multi-run weighting. Mass% is zero when that channel is not selected.</p><table><thead><tr><th>Run</th><th>Initial UV</th><th>Fitted UV</th><th>Initial mass%</th><th>Fitted mass%</th></tr></thead><tbody>{channel_rows}</tbody></table>
{composition_metrics_html}
<h2>Measured and predicted chromatograms</h2><p>The predicted UV and process overlays span each full programmed run. Measured UV is displayed as the median imported sample in each chart pixel column to prevent dark overdraw; every imported sample contributes to the chromatogram objective.</p>{overlays_html}
<h2>Improved shared parameters</h2><p><a href="{FIT_IMPROVED_CSV.name}">Parameter list CSV</a> · <a href="{FIT_IMPROVED_TXT.name}">Parameter list text</a></p><table><thead><tr><th>Parameter</th><th>Initial</th><th>Fitted</th><th>Fit limit ±%</th><th>Allowed min</th><th>Allowed max</th></tr></thead><tbody>{parameter_rows}</tbody></table>
<h2>Profile recipes used</h2>{recipe_details}
<h2>Objective construction</h2><p>UV uses every measured detector sample. RAW_SSE preserves the pointwise UV residual; NORMALIZED_MSE divides by the measured signal scale and sqrt(sample count). Conductivity is included only when enabled. Each entered species mass% is compared with 100 times its integrated detector concentration divided by the integrated concentration of all active species over that collected fraction. CV is cumulative volume, so changing flow is included. The mass% residual group is divided by 100 and sqrt(the populated species-observation count), then expressed in comparable UV units for RAW_SSE. Legacy point-only observations retain midpoint comparison. Blank cells add no residual; chart lines are display-only. Each selected run's process settings remain fixed. Only unlocked active-model parameters are shared and refined.</p>
<h2>Objective improvement over time</h2><p>{history_file_html}</p>{history_plot}
<p>Files: {FIT_TRACE_CSV.name}, {FIT_CONDUCTIVITY_CSV.name}, {FIT_RESULTS_CSV.name}, {FIT_COMPOSITION_CSV.name}, {FIT_SETTINGS_JSON.name}.</p>
""", encoding="utf-8")

    return {
        "success": opt_success,
        "accepted": accepted,
        "acceptance": acceptance,
        "message": opt_message,
        "initial_objective": initial_obj,
        "final_objective": final_obj,
        "improvement_percent": improvement,
        "nfev": opt_nfev,
        "residual_evaluations": eval_count,
        "evaluation_cap_reached": budget_reached,
        "fit_quality_warning": fit_quality_warning,
        "fit_elapsed_seconds": fit_elapsed_seconds,
        "fit_elapsed_hms": fit_elapsed_hms,
        "composition_profile_metrics": composition_profile_metrics,
        "mode": mode,
        "profile_metrics": metric_rows,
        "initial_weighted_rmse": math.sqrt(max(initial_obj, 0.0)) if weighted_objective else None,
        "final_weighted_rmse": math.sqrt(max(final_obj, 0.0)) if weighted_objective else None,
        "improved_parameters_csv": str(FIT_IMPROVED_CSV),
        "improved_parameters_txt": str(FIT_IMPROVED_TXT),
        "fitted_config": fitted,
        "fit_results_csv": str(FIT_RESULTS_CSV),
        "fit_trace_csv": str(FIT_TRACE_CSV),
        "fit_conductivity_csv": str(conductivity_output_path),
        "fit_composition_csv": str(FIT_COMPOSITION_CSV),
        "fit_report_html": str(FIT_REPORT_HTML),
        "fit_config_json": str(FIT_CONFIG_JSON),
        "fit_settings_json": str(FIT_SETTINGS_JSON),
        "fit_history_csv": str(history_path) if history_path else None,
        "latest_fit_history_csv": str(history_path) if history_path else None,
        "apply_jsl": str(FIT_APPLY_JSL),
        "parameters": [spec.path for spec in specs],
    }

"""Integration and regression tests for independent LSQ process setups and weighted RMSE."""
from __future__ import annotations
import copy
import csv
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import sys

import numpy as np
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import tdm_least_squares as ls
from tdm_minimal_io import (
    BATCH_RUNS_CSV, LSQ_RUNS_CSV, apply_batch_shared_parameters,
    batch_row_to_run_config, ensure_lsq_runs_csv, load_batch_rows, load_config,
    load_lsq_rows, save_lsq_rows, validate_config,
)
from tdm_minimal_model import simulate
from tdm_jmp import build_jsl, _validate_jsl
import tdm_bridge
from tdm_minimal_io import MODEL_LABELS, MODEL_EDM_LANGMUIR, MODEL_TDM_LANGMUIR


def config_for(run=1):
    shared = load_config()
    effective, _ = apply_batch_shared_parameters(shared, batch_row_to_run_config(load_lsq_rows()[run-1]))
    effective["numerics"] = {"time_steps": 52, "axial_positions": 8}
    effective["least_squares"]["objective"] = "WEIGHTED_RMSE"
    return effective


def write_excel(path, cv, y, weights):
    wb = Workbook()
    other = wb.active
    other.title = "Readme"
    other.append(["Run signal follows on another worksheet"])
    sheet = wb.create_sheet("Chromatogram")
    sheet.append(["CV", "UV 280 [mAU]", "Weight"])
    for values in zip(cv,y,weights):
        sheet.append([float(v) for v in values])
    wb.save(path)


class RefinementTests(unittest.TestCase):
    def test_jacobian_reuses_current_residual_without_skipping_perturbations(self):
        evaluated = []

        def residual(vector):
            evaluated.append(tuple(vector))
            return np.array([vector[0] + 2 * vector[1], vector[0] - vector[1]])

        cached = ls._cache_last_evaluation(residual)
        x = np.array([1.0, 2.0])
        np.testing.assert_allclose(cached(x), [5.0, -1.0])
        jac = ls._bounded_jacobian(cached, np.zeros(2), np.full(2, 10.0))(x)
        np.testing.assert_allclose(jac, [[1.0, 2.0], [1.0, -1.0]], atol=1e-10)
        self.assertEqual(len(evaluated), 3)  # base + one column solve per parameter

    def test_conductivity_detector_delay_is_not_a_uv_fit_parameter(self):
        cfg = config_for()
        paths = {spec.path for spec in ls.fit_parameter_lock_catalog(cfg)}
        self.assertIn("system.uv_dead_volume_mL", paths)
        self.assertNotIn("system.conductivity_dead_volume_mL", paths)
        trace = simulate(cfg)["trace"]
        changed = copy.deepcopy(cfg)
        changed["system"]["conductivity_dead_volume_mL"] *= 2
        shifted = simulate(changed)["trace"]
        np.testing.assert_allclose(trace["uv_mAU"], shifted["uv_mAU"])

    def test_all_unlocked_edm_fit_reduces_uv_objective(self):
        true = config_for()
        true["numerics"] = {"time_steps": 50, "axial_positions": 6}
        true["least_squares"]["objective"] = "RAW_SSE"
        trace = simulate(true)["trace"]
        initial = copy.deepcopy(true)
        initial["conversion"]["uv_to_protein_mAU_L_g"] *= 0.75
        paths = [spec.path for spec in ls.fit_parameter_lock_catalog(initial)]
        self.assertEqual(len(paths), 10)
        with tempfile.TemporaryDirectory() as d:
            tmp = Path(d)
            source = tmp / "known_uv.csv"
            with source.open("w", newline="") as fh:
                writer = csv.writer(fh)
                writer.writerow(["CV", "UV_mAU"])
                writer.writerows(zip(trace["CV"], trace["uv_mAU"]))
            names = ["FIT_RESULTS_CSV", "FIT_TRACE_CSV", "FIT_COMPOSITION_CSV",
                     "FIT_REPORT_HTML", "FIT_CONFIG_JSON", "FIT_APPLY_JSL",
                     "FIT_SETTINGS_JSON", "FIT_HISTORY_CSV", "FIT_IMPROVED_CSV",
                     "FIT_IMPROVED_TXT", "CONFIG_PATH"]
            with patch.multiple(ls, **{name: tmp / getattr(ls, name).name for name in names}):
                result = ls.run_global_least_squares_refinement(
                    [{"config": initial, "chromatogram_csv": source}],
                    unlocked_paths=paths, max_nfev=4,
                )
        self.assertTrue(result["accepted"])
        self.assertLess(result["final_objective"], result["initial_objective"] * 1e-4)
        self.assertAlmostEqual(result["fitted_config"]["conversion"]["uv_to_protein_mAU_L_g"],
                               true["conversion"]["uv_to_protein_mAU_L_g"], delta=5)
        self.assertEqual(result["parameters"], paths)
        self.assertGreater(result["residual_evaluations"], result["nfev"])

    def test_jmp_capacity_reaches_both_langmuir_solvers(self):
        jsl = build_jsl()
        self.assertIn('Python Name("tdm_ui_c1_qmax_g_L")', jsl)
        for model in (MODEL_LABELS[MODEL_EDM_LANGMUIR], MODEL_LABELS[MODEL_TDM_LANGMUIR]):
            with self.subTest(model=model), patch.dict(tdm_bridge.__dict__, {
                "tdm_ui_model": model,
                "tdm_ui_c1_qmax_g_L": 37.5,
                "tdm_ui_c1_b_L_g": 0.08,
                "tdm_ui_c1_qmax_mg_ml": 999.0,
            }):
                cfg = tdm_bridge.build_shared_config_from_jmp()
                self.assertEqual(cfg["components"][0]["qmax_g_L"], 37.5)
                self.assertAlmostEqual(cfg["components"][0]["H"], 3.0)

    def test_zero_dispersion_fit_starts_at_entered_value(self):
        cfg = config_for()
        cfg["components"][0]["D_app_mm2_s"] = 0.0
        specs = ls.fit_parameter_specs(cfg, unlocked_paths=["components[0].D_app_mm2_s"])
        self.assertEqual(len(specs), 1)
        self.assertEqual(specs[0].initial, 0.0)

    def test_global_fit_can_leave_zero_dispersion_bound(self):
        true = config_for()
        true["numerics"] = {"time_steps": 50, "axial_positions": 6}
        target = float(true["components"][0]["D_app_mm2_s"])
        trace = simulate(true)["trace"]
        base = copy.deepcopy(true)
        base["components"][0]["D_app_mm2_s"] = 0.0
        with tempfile.TemporaryDirectory() as d:
            tmp = Path(d)
            source = tmp / "known_dispersion.csv"
            with source.open("w", newline="") as fh:
                writer = csv.writer(fh)
                writer.writerow(["CV", "UV_mAU"])
                writer.writerows(zip(trace["CV"], trace["uv_mAU"]))
            names = ["FIT_RESULTS_CSV", "FIT_TRACE_CSV", "FIT_COMPOSITION_CSV",
                     "FIT_REPORT_HTML", "FIT_CONFIG_JSON", "FIT_APPLY_JSL",
                     "FIT_SETTINGS_JSON", "FIT_HISTORY_CSV", "FIT_IMPROVED_CSV",
                     "FIT_IMPROVED_TXT", "CONFIG_PATH"]
            with patch.multiple(ls, **{name: tmp / getattr(ls, name).name for name in names}):
                result = ls.run_global_least_squares_refinement(
                    [{"config": base, "chromatogram_csv": source}],
                    unlocked_paths=["components[0].D_app_mm2_s"], max_nfev=20,
                )
        self.assertTrue(result["accepted"])
        self.assertTrue(result["success"], result["message"])
        self.assertLess(result["final_objective"], result["initial_objective"] * 1e-5)
        self.assertAlmostEqual(result["fitted_config"]["components"][0]["D_app_mm2_s"], target, delta=target * .01)

    def test_independent_lsq_process_table(self):
        src = load_batch_rows()
        dst = load_lsq_rows()
        self.assertEqual(len(dst), 30)
        self.assertEqual(dst[0]["LSQ_Weight"], "1.0")
        self.assertEqual(dst[0]["LSQ_X_Origin"], "RUN_START")
        self.assertNotEqual(str(BATCH_RUNS_CSV.resolve()), str(LSQ_RUNS_CSV.resolve()))
        self.assertNotEqual(dst[0]["Run_Name"], src[0]["Run_Name"])

    def test_jsl_includes_new_scope_and_objective(self):
        launcher = (ROOT / "00_RUN_CHROMATOGRAPHY_MODEL_IN_JMP.jsl").read_text(encoding="utf-8")
        self.assertIn('Python Install Packages( "numpy scipy openpyxl" )', launcher)
        self.assertIn('dependencies_r23_checked', launcher)
        jsl = build_jsl()
        _validate_jsl(jsl)
        for expected in ["dtLSQRuns", "SwitchProcessScope", "lsqWeight",
                         "WEIGHTED_RMSE", "lsqRunsFile", "activeRunsFile",
                         "lsqSheet", "lsqXColumn", "lsqSignalColumn", "lsqXUnit", "lsqXOrigin",
                         "Number of runs to fit [1–5]", "Fit settings", "SaveAllLSQReferences"]:
            self.assertIn(expected, jsl)
        self.assertLess(jsl.index("Number of runs to fit [1–5]"), jsl.index("Run section 1"))
        self.assertLess(jsl.index("Run section 5"), jsl.index('Outline Box("Fit settings"'))
        self.assertIn('Column(dtLSQRuns, "LSQ_Chromatogram_CSV")[r] = Trim(Char(LSQPathText(slot)))', jsl)
        self.assertIn('Column(dtLSQRuns, "LSQ_X_Origin")[r] = LSQOriginSelected(slot)', jsl)
        self.assertIn('Python Send(If(actionText == "ATTACH_LSQ_CSV", lsqLoadedRuns[lsqAttachSlot]', jsl)
        self.assertIn('Safety cap: optimizer evaluations', jsl)
        self.assertIn('Refinement running. The JMP window will respond', jsl)
        self.assertNotIn('condDeadVolumeFitLock', jsl)
        self.assertNotIn('LoadLSQReference(1);', jsl)
        self.assertNotIn('lsqRunBoxes[', jsl)
        self.assertNotIn('lsqPathBoxes[', jsl)
        with self.assertRaisesRegex(ValueError, "display boxes cannot"):
            _validate_jsl(jsl.replace("LSQSelectedRun(lsqInitSlot)",
                                      "lsqRunBoxes[lsqInitSlot] << Get Selected"))
        for slot in range(1, 6):
            with self.subTest(slot=slot):
                self.assertIn(f'lsqRunSlot{slot}Row << Visibility(If(Num(lsqRunCount << Get Selected) >= {slot}', jsl)
                self.assertIn(f'Button Box("Attach Excel / CSV", AttachLSQCSV({slot}))', jsl)
                self.assertIn(f'lsqCsvPath{slot} = Text Edit Box', jsl)
                self.assertIn(f'lsqXUnit{slot} = Combo Box', jsl)
                self.assertIn(f'lsqXOrigin{slot} = Combo Box', jsl)
                self.assertIn(f'Return(RunNumberFromLabel(lsqRunSlot{slot} << Get Selected))', jsl)
                self.assertIn(f'lsqCsvPath{slot} << Set Text(path)', jsl)

    def test_conference_cv_and_ml_mau_axes_align_with_solver(self):
        cfg = config_for()
        cfg["column"]["volume_mL"] = 2.5
        cfg["least_squares"]["objective"] = "RAW_SSE"
        trace = simulate(cfg)["trace"]
        program = ls._build_program(cfg, np.zeros(1))
        elution_start = next(program.stage_start_CV[i] for i, stage in enumerate(program.stages)
                             if stage.kind == "ELUTION")
        cv = np.asarray(trace["CV"], dtype=float)
        uv = np.asarray(trace["uv_mAU"], dtype=float)
        keep = cv >= elution_start
        with tempfile.TemporaryDirectory() as d:
            tmp = Path(d)
            whole_run = tmp / "whole_run.xlsx"
            elution_only = tmp / "conference_elution.csv"
            write_excel(whole_run, cv, uv, np.ones(len(cv)))
            with elution_only.open("w", newline="") as fh:
                writer = csv.writer(fh)
                writer.writerow(["Elution Volume (mL)", "Signal", "UV 280 [mAU]"])
                writer.writerows(zip((cv[keep] - elution_start) * 2.5,
                                     np.full(np.count_nonzero(keep), 999.0), uv[keep]))
            info = ls.inspect_chromatogram_csv(elution_only, x_origin="ELUTION_START")
            self.assertEqual((info["x_column"], info["x_unit"], info["signal_column"]),
                             ("Elution Volume (mL)", "ML", "UV 280 [mAU]"))
            full = ls.load_chromatogram_reference(whole_run, cfg)
            partial = ls.load_chromatogram_reference(elution_only, cfg, x_origin="ELUTION_START")
            self.assertTrue(np.allclose(full.cv, cv))
            self.assertTrue(np.allclose(partial.cv, cv[keep]))
            self.assertTrue(np.allclose(partial.signal, uv[keep]))
            residuals, _ = ls._chromatogram_residuals(cfg, partial, trace)
            self.assertLess(float(np.max(np.abs(residuals))), 1e-8)
            with self.assertRaisesRegex(ValueError, "requires CV or mL"):
                ls.inspect_chromatogram_csv(elution_only, x_unit="MIN", x_origin="ELUTION_START")

    def test_xlsx_and_csv_weighted_objective_and_optimisation(self):
        true1 = config_for(1)
        true2 = config_for(2)
        # Change the SECOND LSQ profile recipe so a shared fit truly spans two
        # different chromatographic operating conditions.
        true2["feed"]["total_concentration_mg_mL"] *= 0.8
        true2["column"]["volume_mL"] = 2.5
        # Build a known true UV reference from the existing mechanistic solver.
        sim1 = simulate(true1)["trace"]
        sim2 = simulate(true2)["trace"]
        with tempfile.TemporaryDirectory() as d:
            tmp = Path(d)
            x1, y1 = np.asarray(sim1["CV"]), np.asarray(sim1["uv_mAU"])
            x2, y2 = np.asarray(sim2["CV"]), np.asarray(sim2["uv_mAU"])
            w1 = np.linspace(0.5, 2.0, len(x1))
            w2 = np.ones(len(x2))
            xlsx = tmp/"ref_profile_1.xlsx"
            csvfile = tmp/"ref_profile_2.csv"
            write_excel(xlsx, x1, y1, w1)
            with csvfile.open("w",newline="") as fh:
                writer = csv.writer(fh)
                writer.writerow(["Volume (mL)", "UV 280 [mAU]", "Weight"])
                writer.writerows(zip(x2 * true2["column"]["volume_mL"],y2,w2))
            self.assertEqual(ls.inspect_chromatogram_csv(xlsx)["weight_column"], "Weight")
            r1 = ls.load_chromatogram_reference(xlsx, true1)
            self.assertEqual(len(r1.cv), len(x1))
            self.assertTrue(np.allclose(r1.signal, y1))
            self.assertTrue(np.allclose(r1.point_weights, w1))
            base1 = copy.deepcopy(true1)
            base2 = copy.deepcopy(true2)
            true_uv = float(true1["conversion"]["uv_to_protein_mAU_L_g"])
            base1["conversion"]["uv_to_protein_mAU_L_g"] = true_uv * 0.72
            base2["conversion"]["uv_to_protein_mAU_L_g"] = true_uv * 0.72
            profiles = [
                {"config":base1,"chromatogram_csv":xlsx,"run_context":{"run_number":1,"run_name":"LSQ 1"},"weight":3.0},
                {"config":base2,"chromatogram_csv":csvfile,"run_context":{"run_number":2,"run_name":"LSQ 2"},"weight":1.0},
            ]
            # Prevent test fits from changing the user's original model files.
            path_constants = ["FIT_RESULTS_CSV","FIT_TRACE_CSV","FIT_COMPOSITION_CSV",
                              "FIT_REPORT_HTML","FIT_CONFIG_JSON","FIT_APPLY_JSL",
                              "FIT_SETTINGS_JSON","FIT_HISTORY_CSV","FIT_IMPROVED_CSV",
                              "FIT_IMPROVED_TXT","CONFIG_PATH"]
            with patch.multiple(ls, **{name: tmp / getattr(ls, name).name for name in path_constants}):
                result = ls.run_global_least_squares_refinement(
                    profiles, unlocked_paths=["conversion.uv_to_protein_mAU_L_g"], max_nfev=6)
            self.assertTrue(result["accepted"], result["message"])
            self.assertTrue(result["success"], result["message"])
            self.assertLess(result["final_weighted_rmse"], result["initial_weighted_rmse"] * 0.15)
            fitted_uv = float(result["fitted_config"]["conversion"]["uv_to_protein_mAU_L_g"])
            self.assertAlmostEqual(fitted_uv,true_uv,delta=true_uv*.03)
            self.assertEqual(len(result["profile_metrics"]),2)
            # The objective is literally sum(run normalised weight * weighted point MSE).
            with (tmp/"least_squares_fit_trace.csv").open(newline="") as fh:
                fit_trace = list(csv.DictReader(fh))
            sum_contrib = sum(float(row["Profile_Objective_Contribution"]) for row in fit_trace)
            self.assertAlmostEqual(sum_contrib,result["final_objective"],delta=max(result["final_objective"]*1e-6,1e-8))
            with (tmp/"least_squares_improved_parameters.csv").open(newline="") as fh:
                changed = list(csv.DictReader(fh))
            self.assertEqual([row["Parameter_Path"] for row in changed], ["conversion.uv_to_protein_mAU_L_g"])
            self.assertTrue((tmp/"least_squares_fit_report.html").is_file())
            self.assertTrue((tmp/"least_squares_settings.json").is_file())
            self.assertIn("weighted_rmse_formula",json.loads((tmp/"least_squares_settings.json").read_text()))
            with (tmp/"least_squares_objective_history.csv").open(newline="") as fh:
                history = list(csv.DictReader(fh))
            self.assertEqual(len(history), result["residual_evaluations"] + 1)
            best = [float(row["best_objective"]) for row in history]
            self.assertTrue(all(next_value <= value for value, next_value in zip(best, best[1:])))

    def test_xlsx_column_override_and_independent_metadata(self):
        from tdm_bridge import _lsq_reference_column_options
        with tempfile.TemporaryDirectory() as d:
            source = Path(d)/"multiple_detector_signals.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Baseline data"
            sheet.append(["CV", "Signal"])
            for n in range(6):
                sheet.append([n, 10.0 + n])
            signal = workbook.create_sheet("Actual target")
            signal.append(["Sample index", "UV 280 [mAU]", "Measured response", "Point Weight"])
            for n in range(6):
                signal.append([n, 99, 10.0 + 2*n, 1.0 + n])
            workbook.save(source)
            row = {"LSQ_Sheet": "Actual target", "LSQ_X_Column": "Sample index",
                   "LSQ_Signal_Column": "Measured response", "LSQ_X_Unit": "INDEX"}
            options = _lsq_reference_column_options(row)
            self.assertEqual(options["x_origin"], "RUN_START")
            info = ls.inspect_chromatogram_csv(source, **options)
            self.assertEqual(info["signal_column"], "Measured response")
            self.assertEqual(info["weight_column"], "Point Weight")
            cfg = config_for()
            ref = ls.load_chromatogram_reference(source, cfg, **options)
            self.assertEqual(ref.x_unit, "INDEX")
            self.assertEqual(ref.signal.tolist(), [10,12,14,16,18,20])
            self.assertEqual(ref.point_weights.tolist(), [1,2,3,4,5,6])
            with self.assertRaisesRegex(ValueError, "worksheet"):
                ls.inspect_chromatogram_csv(source, sheet_name="Missing tab")
            with self.assertRaisesRegex(ValueError, "also set the reference X-axis unit"):
                ls.inspect_chromatogram_csv(source, sheet_name="Actual target", x_column="UV 280 [mAU]", signal_column="Measured response", x_unit="AUTO")

    def test_weighted_rmse_rejects_composition_mix(self):
        # A composition mass% residual must not silently distort a response-unit RMSE.
        with tempfile.TemporaryDirectory() as d:
            source = Path(d)/"chrom.csv"
            source.write_text("CV,UV 280 [mAU]\n0,1\n1,2\n2,3\n", encoding="utf-8")
            cfg = config_for()
            with self.assertRaisesRegex(ValueError, "chromatogram signals only"):
                ls.run_global_least_squares_refinement([{
                    "config": cfg, "chromatogram_csv": source, "composition_references": [
                        {"CV": 0.5, "mass_percent": {"Target": 100.0}}],
                }], mode="CHROMATOGRAM_AND_COMPOSITION", unlocked_paths=[])

    def test_invalid_sample_and_run_weights_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            file = Path(d)/"broken.csv"
            file.write_text("CV,Signal,Weight\n0,1,1\n1,2,-1\n2,3,1\n",encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "nonnegative"):
                ls.load_chromatogram_reference(file, config_for())
            file.write_text("CV,Signal\n0,1\n1,2\n2,3\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "run weight"):
                ls.run_global_least_squares_refinement([{
                    "config": config_for(), "chromatogram_csv": file, "weight": 0,
                }], unlocked_paths=[])
            with self.assertRaisesRegex(ValueError, "1–5"):
                ls.run_global_least_squares_refinement([{}] * 6)


if __name__ == "__main__":
    unittest.main(verbosity=2)
